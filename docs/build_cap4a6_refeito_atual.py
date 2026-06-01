from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
IMAGES = DOCS / "imagens"
OUT = DOCS / "Relatorio_Rui_Agostinho_cap4a6_refeito_31_05_2026.docx"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text: str, bold: bool = False, size: float = 9.5) -> None:
    cell.text = ""
    paragraph = cell.paragraphs[0]
    paragraph.paragraph_format.space_after = Pt(0)
    run = paragraph.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    run.font.name = "Calibri"
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table) -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = "w:{}".format(edge)
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "4")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "DADCE0")


def style_document(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.10

    for name, size, color, before, after in [
        ("Heading 1", 16, "2E74B5", 16, 8),
        ("Heading 2", 13, "2E74B5", 12, 6),
        ("Heading 3", 12, "1F4D78", 8, 4),
    ]:
        style = styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    if "CodeBlock" not in styles:
        code = styles.add_style("CodeBlock", 1)
    else:
        code = styles["CodeBlock"]
    code.font.name = "Courier New"
    code.font.size = Pt(8)
    code.paragraph_format.left_indent = Inches(0.15)
    code.paragraph_format.right_indent = Inches(0.05)
    code.paragraph_format.space_after = Pt(0)
    code.paragraph_format.line_spacing = 1.0

    if "CaptionLocal" not in styles:
        caption = styles.add_style("CaptionLocal", 1)
    else:
        caption = styles["CaptionLocal"]
    caption.font.name = "Calibri"
    caption.font.size = Pt(9)
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(79, 79, 79)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(8)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run("Virtual CISO - capítulos 4 a 6 revistos").font.size = Pt(8)


def p(doc: Document, text: str, style: str | None = None, bold_prefix: str | None = None) -> None:
    paragraph = doc.add_paragraph(style=style)
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        run.bold = True
        paragraph.add_run(text[len(bold_prefix):])
    else:
        paragraph.add_run(text)


def h(doc: Document, level: int, text: str) -> None:
    doc.add_heading(text, level=level)


def bullet(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Bullet")


def numbered(doc: Document, text: str) -> None:
    doc.add_paragraph(text, style="List Number")


def table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float] | None = None) -> None:
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_table_borders(t)
    for i, header in enumerate(headers):
        set_cell_text(t.rows[0].cells[i], header, bold=True, size=9)
        set_cell_shading(t.rows[0].cells[i], "F2F4F7")
    for row in rows:
        cells = t.add_row().cells
        for i, value in enumerate(row):
            set_cell_text(cells[i], value, size=8.8)
    if widths:
        for row in t.rows:
            for i, width in enumerate(widths):
                row.cells[i].width = Inches(width)
    doc.add_paragraph()


def figure(doc: Document, filename: str, caption: str) -> None:
    path = IMAGES / filename
    if not path.exists():
        p(doc, f"[Figura não inserida: {path}]")
        return
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(path), width=Inches(6.25))
    cap = doc.add_paragraph(caption, style="CaptionLocal")
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER


def source_path(relative: str) -> Path:
    return ROOT / relative


def code_snippet(doc: Document, title: str, relative: str, start: int, end: int, why: str) -> None:
    h(doc, 2, title)
    p(doc, f"Localização: {relative} (linhas {start}-{end}).")
    p(doc, why)
    path = source_path(relative)
    if not path.exists():
        p(doc, f"[Ficheiro não encontrado: {relative}]")
        return
    lines = path.read_text(encoding="utf-8", errors="replace").splitlines()
    selected = lines[start - 1:end]
    for line_no, line in enumerate(selected, start=start):
        cp = doc.add_paragraph(style="CodeBlock")
        cp.add_run(f"{line_no:>4}  {line}")


def add_cover(doc: Document) -> None:
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Virtual CISO")
    run.bold = True
    run.font.size = Pt(22)
    run.font.color.rgb = RGBColor.from_string("0B2545")

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = subtitle.add_run("Capítulos 4, 5 e 6 revistos com evidência da implementação")
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor.from_string("333333")

    p(doc, "Documento de trabalho para substituição dos capítulos de conceção, demonstração e conclusões no relatório principal. Foi preparado a partir da estrutura do relatório de 18/05/2026, mas atualizado com as capacidades efetivamente implementadas na aplicação em 31/05/2026.")
    p(doc, "As capturas de ecrã foram recolhidas da instância local autenticada da aplicação e guardadas em docs/imagens/. Os anexos incluem excertos de código e a localização dos ficheiros no repositório, de modo a tornar a implementação demonstrável e auditável.")
    doc.add_page_break()


def chapter_4(doc: Document) -> None:
    h(doc, 1, "4. Conceção e Desenvolvimento do Virtual CISO")
    p(doc, "Este capítulo descreve a conceção, a arquitetura e a implementação do artefacto desenvolvido, designado Virtual CISO. A exposição segue a lógica de Design Science Research, em que a relevância do artefacto depende da ligação explícita entre problema, requisitos, construção e avaliação (Hevner et al., 2004; Peffers et al., 2007). Ao contrário de uma descrição meramente tecnológica, o capítulo procura demonstrar como cada decisão de implementação responde às lacunas identificadas na revisão sistemática: fragmentação entre GRC, dashboards centrados em reporte, ausência de conformidade contínua, fraca rastreabilidade multi-framework e uso pouco governado de IA.")
    p(doc, "A versão atual da aplicação já não deve ser entendida como um conjunto de páginas isoladas. O frontend foi consolidado em torno de uma home programática para o CISO, que calcula a etapa atual do programa, a ação seguinte e a justificação operacional com base nos dados reais da organização. Esta alteração é central para a tese: transforma uma arquitetura integrada ao nível dos dados num mecanismo integrado ao nível do trabalho humano.")

    h(doc, 2, "4.1. Requisitos do Sistema")
    p(doc, "Os requisitos do Virtual CISO foram derivados de três fontes: as lacunas identificadas na literatura, o enquadramento regulatório aplicável a organizações sujeitas a NIS2 e a análise prática das responsabilidades do CISO. A formulação dos requisitos procurou evitar que o artefacto se limitasse a agregar dados. O objetivo foi construir um sistema capaz de orientar decisões, justificar prioridades e preservar rastreabilidade auditável.")
    table(doc, ["Lacuna", "Resposta implementada", "Evidência no artefacto"], [
        ["Fragmentação GRC", "Modelo integrado de ativos, riscos, controlos, mecanismos, evidências e frameworks.", "Serviços de rastreabilidade, propagação de conformidade e home do programa do CISO."],
        ["Reporte sem decisão", "Fila de próxima-melhor-ação por etapa do programa.", "Mission Control em modo programático e ProgramService."],
        ["Conformidade pontual", "Snapshots de postura e deteção de drift.", "Página Regressão de postura e drift e SecurityPostureDriftService."],
        ["Avaliação pouco comparável", "Pontuação por framework e rastreabilidade L4.", "Explorador de rastreabilidade e cadeia controlo externo -> controlo interno -> mecanismo -> evidência."],
        ["IA sem governação", "RAG com fontes clicáveis, modelo ponderado explicável e portão XGBoost+SHAP.", "Assistente CISO, KnowledgeSource, PriorityModelConfig e feature store."],
    ], widths=[1.45, 2.55, 2.5])

    h(doc, 2, "4.2. Arquitetura Geral do Sistema")
    p(doc, "A arquitetura do Virtual CISO segue uma organização modular em camadas. A camada de apresentação é uma aplicação React com TypeScript, responsável pela experiência de navegação, dashboards, exploradores e fluxos de decisão. A camada de serviços é implementada em Django REST Framework e concentra a lógica de domínio: risco, conformidade, governação, rastreabilidade, assistente e integrações externas. A camada de persistência assenta em PostgreSQL, incluindo armazenamento relacional e suporte a embeddings através de pgvector para o RAG semântico.")
    p(doc, "A camada de integração agrega fontes externas de segurança, incluindo NVD/NIST, FIRST EPSS, CISA KEV, Wazuh e Nmap. Esta camada não é periférica: alimenta diretamente a priorização contextual, a entrada de ativos e a avaliação contínua da postura. A camada analítica combina regras determinísticas, modelos ponderados explicáveis e, em modo governado, um estimador interno XGBoost+SHAP preparado para ativação futura. Por fim, a camada de IA generativa opera como interface explicativa e de síntese, e não como fonte autónoma de decisão.")
    figure(doc, "fig-4-01-ciclo-programa-ciso.png", "Figura 4.1 - Home programática do CISO, mostrando foco atual, progresso por etapas e sinais operacionais calculados a partir dos dados reais.")

    h(doc, 2, "4.3. Implementação Técnica")
    p(doc, "O backend organiza-se por domínios Django, designadamente risk, governance, ciso_assistant e integrations. Esta separação permite manter a lógica de risco, governação e IA em serviços próprios, reduzindo o acoplamento entre endpoints e regras de negócio. No frontend, as páginas foram reorganizadas para reduzir duplicações na navegação e tornar as integrações externas visíveis como tal. A CISA KEV, por exemplo, deixou de ser um botão solto no dicionário de vulnerabilidades e passou a estar integrada no grupo de Integrações.")
    p(doc, "A implementação técnica privilegia serviços de domínio reutilizáveis. Em vez de replicar regras no frontend, a aplicação expõe endpoints de overview, traceability, drift e priority-model-config. Esta opção reduz divergências entre apresentação e domínio e permite que a demonstração seja sustentada por código verificável. O Anexo C identifica os ficheiros e excertos mais relevantes.")

    h(doc, 2, "4.4. Modelo de Dados e Integração de Informação")
    p(doc, "O modelo de dados integra entidades técnicas e organizacionais. No domínio de risco, os ativos são enriquecidos com tipo, responsável, criticidade, exposição, software e vulnerabilidades. No domínio de governação, as entidades centrais são frameworks, controlos externos, controlos internos, mecanismos, evidências, avaliações e ações. O ponto essencial é que estas entidades não vivem em silos: a priorização de vulnerabilidades consulta o contexto organizacional do ativo; a conformidade usa evidências e mecanismos; a rastreabilidade percorre a cadeia completa entre requisito normativo e evidência concreta.")
    p(doc, "A entrada de ativos foi redesenhada com uma regra de qualidade explícita: ativos descobertos por Wazuh ou Nmap não entram no inventário oficial sem responsável de negócio, tipo de ativo e classificação validada com justificação auditável. Esta regra resolve uma fragilidade frequente em inventários técnicos: a existência de ativos detetados não equivale a informação governável. O inventário oficial passa, assim, a representar apenas ativos suficientemente caracterizados para suportar decisões de risco e conformidade.")
    figure(doc, "fig-4-02-onboarding-ativos.png", "Figura 4.2 - Entrada de ativos com backlog de onboarding, métricas acionáveis e regra de validação antes de entrada no inventário oficial.")

    h(doc, 2, "4.5. Mecanismos de Implementação e Evidência")
    p(doc, "Um contributo estrutural do Virtual CISO é a distinção entre requisito, controlo e mecanismo. Um requisito normativo define o que deve ser assegurado; um controlo interno expressa a forma organizacional adotada; um mecanismo representa a implementação concreta; e a evidência demonstra a sua existência, validade ou eficácia. Esta camada intermédia evita que a conformidade seja tratada como uma marca binária por framework, permitindo avaliar se a organização tem mecanismos reais e evidência reutilizável.")
    p(doc, "A evidência é tratada como entidade auditável e reutilizável. Quando associada a mecanismos, controlos ou documentos, pode propagar impacto para várias frameworks sem duplicação. Esta opção é particularmente relevante em ambientes multi-framework, nos quais ISO/IEC 27001, NIS2, NIST CSF ou QNRC podem partilhar controlos equivalentes.")

    h(doc, 2, "4.6. Gestão de Risco e Priorização Contextual")
    p(doc, "A priorização de vulnerabilidades responde diretamente à RQ2. O sistema não ordena ocorrências apenas por CVSS; combina severidade técnica, probabilidade de exploração EPSS, presença no catálogo CISA KEV, criticidade do ativo, exposição, valor de negócio, dependências, lacunas de mecanismos, lacunas de evidência, relevância regulatória e diferença de risco residual. O resultado é um score multidimensional acompanhado por decomposição de contribuições.")
    p(doc, "O EPSS é usado como sinal preditivo externo de exploração, produzido pela FIRST através de modelos de machine learning. A CISA KEV é tratada como evidência de exploração conhecida, reforçando a prioridade quando uma vulnerabilidade consta do catálogo Known Exploited Vulnerabilities. Na instância analisada existem 1159 CVEs no catálogo, dos quais 13 estão marcados como presentes em CISA KEV. Esta integração está exposta numa página própria de administração, coerente com o seu estatuto de fonte externa.")
    figure(doc, "fig-5-03-cisa-kev-integracao.png", "Figura 4.3 - Integração CISA KEV com métricas de enriquecimento e estado de sincronização.")

    h(doc, 3, "4.6.1. Modelo ponderado explicável")
    p(doc, "O modelo atualmente servido em produção é ponderado e explicável. Para cada fator, o sistema guarda valor bruto, score normalizado, peso e contribuição. Esta decomposição é importante para defesa e auditoria: permite explicar porque uma vulnerabilidade se encontra acima de outra mesmo quando a pontuação CVSS é idêntica. O modelo é determinístico e, portanto, reprodutível.")

    h(doc, 3, "4.6.2. Evolução governada para XGBoost+SHAP")
    p(doc, "A aplicação está preparada para ativar um estimador interno XGBoost+SHAP, mas a ativação é bloqueada enquanto não existir histórico temporal suficiente. Esta distinção é essencial na dissertação: o sinal de machine learning atualmente ativo é o EPSS externo; o XGBoost+SHAP interno encontra-se implementado, instrumentado e governado, mas condicionado por dados organizacionais suficientes. Na instância atual, o feature store contém 1176 snapshots de features, 195 de 200 desfechos rotulados exigidos e 1 de 2 ciclos temporais exigidos. As dependências xgboost, shap, sklearn, numpy e joblib estão instaladas, mas ainda faltam artefacto treinado e métricas holdout.")
    figure(doc, "fig-5-02-modelo-priorizacao.png", "Figura 4.4 - Painel de administração do modelo de priorização, com portão de prontidão para XGBoost+SHAP.")

    h(doc, 2, "4.7. Conformidade Multi-Framework e Avaliação Transversal")
    p(doc, "A conformidade é calculada a partir de controlos, mecanismos e evidências, e não declarada manualmente como estado isolado. A aplicação suporta avaliação transversal por framework e permite observar frameworks com cobertura insuficiente, controlos sem avaliação e gaps ainda por tratar. Esta opção responde à lacuna de avaliação comparável por referencial, pois torna possível comparar posture scores mantendo a rastreabilidade dos elementos que os sustentam.")
    p(doc, "A avaliação foi desenhada de forma transversal: a etapa Avaliar do programa do CISO não é específica do CNCS, mas aplica-se a todas as frameworks em âmbito. Assim, NIS2, ISO/IEC 27001, NIST CSF, QNRC ou DL 125/2025 são tratados segundo a mesma lógica de postura, cobertura e lacunas.")

    h(doc, 2, "4.8. Rastreabilidade Transversal")
    p(doc, "O explorador de rastreabilidade é a camada que materializa a cadeia ativo -> risco -> controlo -> mecanismo -> evidência -> framework. A versão atual permite re-enraizar a exploração dentro da própria página: ao clicar num mecanismo, controlo, evidência ou framework, o utilizador não abandona o contexto, mas passa a explorar a cadeia a partir desse novo ponto. Isto transforma a rastreabilidade de uma visualização estática numa ferramenta de auditoria navegável.")
    p(doc, "A rastreabilidade suporta modos oficial, simulação e exploratório. Esta distinção evita misturar evidência aprovada com evidência em revisão quando se calcula a posição oficial, mas permite ao CISO simular o impacto de validações futuras. Na captura abaixo, uma framework apresenta 506 cadeias L4, ligando controlos externos a controlos internos, mecanismos e evidência.")
    figure(doc, "fig-5-04-rastreabilidade-framework.png", "Figura 4.5 - Explorador de rastreabilidade com cadeia L4 entre controlo externo, controlo interno, mecanismos e evidência.")

    h(doc, 2, "4.9. IA Generativa e RAG Híbrido")
    p(doc, "O assistente de IA adota uma arquitetura RAG híbrida, combinando recuperação semântica, recuperação estruturada e resposta em linguagem natural. A base RAG contém 4776 chunks indexados, incluindo ativos, vulnerabilidades, controlos, mecanismos, evidências, documentos de governação, lacunas de conformidade e regulamentos técnicos. Para fontes normativas como NIS2 e DL 125/2025, a ingestão preserva metadados de artigo, parágrafo, página e ficheiro de origem.")
    p(doc, "A arquitetura segue o princípio internal-first: quando a pergunta diz respeito ao estado da organização, o sistema privilegia dados internos de governação e evidência antes de recorrer a conhecimento normativo externo. As respostas apresentam fontes rastreáveis e clicáveis. Quando o utilizador abre uma fonte, a aplicação mostra o documento, artigo, parágrafo, página e excerto citado, reforçando a verificabilidade das respostas.")
    figure(doc, "fig-5-05-fonte-nis2-rag.png", "Figura 4.6 - Fonte normativa NIS2 aberta a partir do mecanismo de citações RAG, com artigo, parágrafo, páginas e excerto citado.")

    h(doc, 2, "4.10. Conformidade Contínua e Drift")
    p(doc, "A conformidade contínua é suportada por snapshots persistidos de avaliações de controlos e por comparação temporal da exposição técnica dos ativos. O serviço de drift é deliberadamente determinístico: o LLM pode ser usado para sintetizar recomendações, mas a prova de regressão vem da base de dados e dos snapshots. Esta separação reforça a auditabilidade do mecanismo.")
    p(doc, "Na instância atual existe uma regressão crítica de controlo detetada face ao snapshot de baseline. O painel apresenta baseline, número de snapshots, métricas de regressão e recomendações. Esta funcionalidade responde diretamente à lacuna de auditorias pontuais, deslocando a conformidade para um modelo de monitorização contínua.")
    figure(doc, "fig-5-06-drift-conformidade.png", "Figura 4.7 - Página de regressão de postura e drift, com baseline, eventos críticos e passos de revisão.")

    h(doc, 2, "4.11. Fio Condutor do CISO")
    p(doc, "A consolidação mais relevante do frontend é a introdução do fio condutor de trabalho do CISO. Em vez de apresentar dashboards concorrentes, a home do programa calcula oito etapas: Contexto, Conhecer, Avaliar, Priorizar, Tratar, Provar, Vigiar e Reportar. Cada etapa tem completude, estado, bloqueios e ações associadas. A etapa atual é escolhida a partir dos dados reais e não por uma checklist decorativa.")
    p(doc, "Na captura recolhida, o foco atual é Conhecer, com 53 ativos por integrar e ação seguinte orientada para ativos sem classificação validada. Isto demonstra uma passagem importante: a aplicação não se limita a dizer que há problemas; calcula a próxima ação que desbloqueia o programa. Esta capacidade liga a lacuna transversal da SLR à experiência real de utilização.")

    h(doc, 2, "4.12. Mapa de Implementação")
    p(doc, "Para reforçar a verificabilidade técnica, a Tabela 4.2 sintetiza a correspondência entre capacidades descritas e ficheiros do repositório. Os excertos selecionados encontram-se no Anexo C.")
    table(doc, ["Capacidade", "Backend", "Frontend"], [
        ["Fio condutor do CISO", "backend/governance/services/program_service.py; backend/governance/views_program.py", "frontend/src/pages/MissionControl.tsx"],
        ["Onboarding e inventário fiável", "backend/risk/views.py; backend/risk/serializers.py", "frontend/src/pages/assets/AssetOnboarding.tsx; frontend/src/pages/assets/Inventory.tsx"],
        ["Priorização contextual", "backend/risk/services/prioritization.py; backend/risk/models/priority_model.py", "frontend/src/pages/risk/RiskPrioritization.tsx; frontend/src/pages/admin/PriorityModel.tsx"],
        ["CISA KEV", "backend/risk/services/intel_service.py; backend/integrations/views.py", "frontend/src/pages/admin/integrations/KEV.tsx; frontend/src/pages/Vulnerabilities.tsx"],
        ["RAG com citações", "backend/ciso_assistant/management/commands/ingest_nis2_normative.py; backend/ciso_assistant/views.py", "frontend/src/pages/Assistant.tsx; frontend/src/pages/KnowledgeSource.tsx"],
        ["Rastreabilidade", "backend/governance/services/traceability_service.py", "frontend/src/pages/governance/TraceabilityExplorer.tsx"],
        ["Drift", "backend/governance/services/security_posture_drift.py", "frontend/src/pages/governance/ComplianceDrift.tsx"],
    ], widths=[1.65, 2.65, 2.2])


def chapter_5(doc: Document) -> None:
    h(doc, 1, "5. Demonstração e Avaliação do Virtual CISO")
    p(doc, "Este capítulo demonstra e avalia o artefacto implementado, com foco em três casos de uso alinhados com as questões de investigação e com as lacunas identificadas na revisão sistemática. A avaliação assume a natureza de prova de conceito própria da DSR: procura demonstrar utilidade, adequação e rastreabilidade técnica, sem reivindicar validação estatística em ambiente produtivo.")

    h(doc, 2, "5.1. Cenário de Demonstração")
    p(doc, "A demonstração usa a instância local do Virtual CISO, com dados reais de desenvolvimento e cenários controlados. No momento da recolha, a base continha 54 ativos, dos quais 53 permaneciam em onboarding, 1159 CVEs no catálogo, 13 CVEs em CISA KEV, 1176 snapshots de features de priorização, 4776 chunks RAG e uma regressão crítica de postura detetada por drift. Estes valores não são exemplificativos: foram extraídos da base da aplicação e usados para orientar as capturas.")
    table(doc, ["Dimensão", "Valor observado", "Interpretação"], [
        ["Ativos", "54 total; 53 em onboarding", "A organização ainda está na etapa Conhecer, pois a maioria dos ativos carece de responsável, tipo ou classificação."],
        ["Vulnerabilidades", "1159 CVEs; 13 em CISA KEV", "A priorização tem volume suficiente para exigir ordenação contextual."],
        ["Modelo XGBoost+SHAP", "1176 snapshots; 195/200 desfechos; 1/2 ciclos", "O estimador interno está preparado mas bloqueado por prontidão, mantendo o modelo ponderado oficial."],
        ["RAG", "4776 chunks; fonte NIS2 art. 21.º disponível", "A pergunta normativa pode ser respondida com citação ao artigo/parágrafo."],
        ["Drift", "1 regressão crítica", "A conformidade contínua está operacional e evidencia degradação face ao snapshot."],
    ], widths=[1.45, 1.6, 3.45])

    h(doc, 2, "5.2. UC1 - Priorização de Vulnerabilidades Sensível ao Contexto Organizacional")
    p(doc, "O primeiro caso de uso responde à pergunta do CISO: 'Apareceram vulnerabilidades novas; por onde começo?'. A demonstração utiliza duas vulnerabilidades com CVSS idêntico, mas contexto organizacional distinto. A primeira afeta o Portal de Munícipes, ativo crítico e exposto, com EPSS elevado e presença em CISA KEV. A segunda afeta um servidor QA interno, com EPSS reduzido e mitigação acionável. O sistema classifica a primeira em #1 com prioridade 73 e a segunda em #2 com prioridade 39 no conjunto demonstrativo.")
    table(doc, ["Critério", "DEMO UC1 - Portal de Munícipes", "DEMO UC1 - Servidor QA Interno"], [
        ["CVSS", "8.0", "8.0"],
        ["EPSS", "78%", "5%"],
        ["CISA KEV", "Sim", "Não"],
        ["Exposição", "Elevada / público", "Interna / controlada"],
        ["Resultado", "#1; prioridade 73", "#2; prioridade 39"],
        ["Explicação", "CVSS, criticidade, EPSS, KEV e exposição aumentam o score.", "Mitigação e menor exposição reduzem a prioridade."],
    ], widths=[1.25, 2.6, 2.65])
    figure(doc, "fig-5-01-priorizacao-contextual.png", "Figura 5.1 - Use Case 1: ranking contextual com duas vulnerabilidades de CVSS idêntico mas prioridade distinta.")
    p(doc, "A importância desta demonstração reside na explicabilidade. O CISO não recebe apenas uma lista ordenada: recebe o motivo. A decomposição por fator permite justificar à gestão de topo que uma vulnerabilidade tecnicamente equivalente deve ser tratada antes por afetar um ativo crítico exposto, explorável e sem mitigação suficiente.")

    h(doc, 2, "5.3. UC2 - Q&A sobre Conformidade NIS2/ISO 27001 com Rastreabilidade")
    p(doc, "O segundo caso de uso responde à pergunta: 'Estamos conformes com o artigo 21.º da NIS2 sobre gestão de riscos?'. A aplicação indexa a NIS2 como fonte normativa com metadados de artigo, parágrafo e página, permitindo que o assistente recupere o artigo aplicável e o relacione com controlos, mecanismos, evidências e lacunas internos. A fonte aberta na Figura 5.2 mostra o Artigo 21.º, n.º 1, com páginas 48-49 e excerto citado.")
    figure(doc, "fig-5-05-fonte-nis2-rag.png", "Figura 5.2 - Fonte normativa NIS2 rastreável ao artigo, parágrafo, páginas e ficheiro PDF.")
    p(doc, "Esta capacidade responde à lacuna de conformidade contínua e evidência testável, pois a resposta do assistente não depende apenas de texto gerado. Cada afirmação pode remeter para fonte normativa ou entidade interna, permitindo ao CISO abrir o documento de origem e verificar a base da resposta.")

    h(doc, 2, "5.4. UC3 - Deteção de Regressão de Conformidade")
    p(doc, "O terceiro caso de uso responde à pergunta: 'Algo na nossa postura de segurança piorou desde a última auditoria?'. A funcionalidade de drift cria ou utiliza snapshots de avaliações de controlos e compara o estado atual com a fotografia anterior. Quando existe degradação, a página apresenta o evento, a severidade, a baseline e uma recomendação de revisão. Na instância atual foi detetada uma regressão crítica de controlo.")
    figure(doc, "fig-5-06-drift-conformidade.png", "Figura 5.3 - Use Case 3: deteção de regressão de postura face a snapshot anterior.")
    p(doc, "A demonstração mostra que o sistema não fica limitado a auditorias pontuais. A possibilidade de criar snapshots e comparar a evolução aproxima a conformidade de um processo contínuo, essencial em contextos regulados por NIS2.")

    h(doc, 2, "5.5. Avaliação da Rastreabilidade e da Postura por Framework")
    p(doc, "A rastreabilidade foi avaliada através do explorador transversal. Para uma framework selecionada, o sistema apresenta score oficial, score de simulação, controlos internos, mecanismos, evidências e cadeias L4. A presença de 506 cadeias na framework DL 125/2025 evidencia que a avaliação por framework não é apenas uma grelha de catálogo: é navegável até mecanismos e evidência.")
    figure(doc, "fig-5-04-rastreabilidade-framework.png", "Figura 5.4 - Cadeia L4 no explorador de rastreabilidade.")
    p(doc, "A funcionalidade tem valor particular para a defesa da dissertação, porque torna visual o fio condutor ativo/controlo/mecanismo/evidência/framework. Esta cadeia é a materialização prática da lacuna transversal identificada na literatura.")

    h(doc, 2, "5.6. Avaliação do Fio Condutor do CISO")
    p(doc, "A home do programa foi avaliada enquanto mecanismo de orquestração. A aplicação calcula a etapa atual e a próxima ação a partir dos sinais de ativos, risco, maturidade, drift e governance work items. No estado observado, a etapa atual é Conhecer, com a ação seguinte 'Ativos sem classificação validada'. Isto demonstra que o sistema passou de dashboard para assistente de trabalho: não apenas mostra indicadores, mas indica a sequência operacional mais relevante.")
    figure(doc, "fig-4-01-ciclo-programa-ciso.png", "Figura 5.5 - Fio condutor do CISO com foco atual, etapas concluídas e sinais de bloqueio.")

    h(doc, 2, "5.7. Avaliação do Onboarding e Inventário Fiável")
    p(doc, "A entrada de ativos foi avaliada pela regra de passagem para inventário oficial. A página de onboarding apresenta backlog, filtros acionáveis, seleção múltipla e triagem. A regra implementada impede que ativos ou software descobertos automaticamente apareçam como inventário finalizado sem responsável, tipo e classificação validada. Esta decisão reduz o risco de a aplicação apresentar informação técnica incompleta como se fosse base de governação fiável.")

    h(doc, 2, "5.8. Avaliação do Modelo de Priorização e do Portão XGBoost+SHAP")
    p(doc, "A página de administração do modelo mostra que a aplicação distingue claramente o modelo oficial ponderado do estimador interno XGBoost+SHAP. O modo XGBoost+SHAP está bloqueado até cumprir prontidão: desfechos rotulados, ciclos temporais, artefacto treinado e métricas de validação. Esta escolha é metodologicamente relevante: evita afirmar que existe um modelo organizacional treinado quando os dados ainda não são suficientes, mas demonstra que a arquitetura já recolhe o dataset necessário.")

    h(doc, 2, "5.9. Síntese de Avaliação por Requisitos")
    table(doc, ["Requisito", "Estado", "Evidência"], [
        ["Integração GRC", "Cumprido", "Programa do CISO, rastreabilidade e serviços de overview agregam risco, ativos, conformidade e evidências."],
        ["Priorização contextual", "Cumprido", "UC1 demonstra CVSS igual com prioridades diferentes por contexto."],
        ["RAG rastreável", "Cumprido", "Fonte NIS2 aberta por artigo/parágrafo/página e chunks RAG indexados."],
        ["Conformidade contínua", "Cumprido em prova de conceito", "Snapshots e deteção de regressão crítica na página Drift."],
        ["XGBoost+SHAP interno", "Implementado mas não ativado", "Portão bloqueia ativação até cumprir histórico e métricas."],
        ["Inventário fiável", "Cumprido", "Ativos só entram no inventário final após responsável, tipo e classificação validada."],
    ], widths=[1.7, 1.55, 3.25])

    h(doc, 2, "5.10. Discussão Face às Lacunas e Questões de Investigação")
    p(doc, "Os resultados permitem afirmar que o artefacto responde às lacunas principais identificadas na SLR. A fragmentação é mitigada pela agregação programática e pela rastreabilidade transversal. A crítica aos dashboards centrados em reporte é respondida pela próxima-melhor-ação e pelo fluxo do CISO. A conformidade contínua é abordada pelo drift. A avaliação comparável por framework é suportada por scores e cadeias L4. Finalmente, a IA é usada de forma governada: EPSS como sinal preditivo externo, RAG com fontes e XGBoost+SHAP interno condicionado por prontidão.")
    p(doc, "A avaliação deve, contudo, ser interpretada como demonstração de utilidade de um artefacto e não como estudo longitudinal em produção. A existência de cenários controlados é adequada à fase DSR, mas não substitui validação posterior com CISOs reais, organizações reguladas ou auditorias independentes.")

    h(doc, 2, "5.11. Ameaças à Validade")
    bullet(doc, "Validade interna: alguns casos usam dados de demonstração controlada, pelo que os valores não provam desempenho estatístico generalizável.")
    bullet(doc, "Validade externa: a instância local simula uma organização, mas não substitui operação prolongada num contexto produtivo.")
    bullet(doc, "Validade de construção: os indicadores de completude e prontidão dependem de limiares definidos pelo autor; por isso, devem ser documentados e, futuramente, parametrizados pelo apetite de risco.")
    bullet(doc, "Validade técnica: a dependência de LLM local e embeddings pode introduzir variação de latência e qualidade, embora as decisões críticas permaneçam determinísticas.")


def chapter_6(doc: Document) -> None:
    h(doc, 1, "6. Conclusões e Trabalho Futuro")
    p(doc, "A dissertação teve como objetivo conceber, implementar e avaliar um artefacto de apoio à decisão para o CISO, capaz de integrar governação, risco, conformidade e IA de forma explicável e rastreável. O Virtual CISO resultante não deve ser entendido como uma substituição do decisor humano, mas como uma plataforma de orientação, síntese e justificação.")

    h(doc, 2, "6.1. Síntese do Trabalho Desenvolvido")
    p(doc, "O trabalho começou pela identificação de lacunas na literatura e evoluiu para a construção de um protótipo funcional. A aplicação integra ativos, vulnerabilidades, controlos, mecanismos, evidências, frameworks e fontes normativas. A nível de interface, a evolução mais relevante foi a passagem de dashboards concorrentes para uma home programática que orienta o CISO por etapas. A nível analítico, a priorização contextual demonstra que risco organizacional não é equivalente a severidade técnica. A nível de conformidade, a rastreabilidade e o drift transformam evidência e auditoria em processos contínuos.")

    h(doc, 2, "6.2. Resposta às Questões de Investigação")
    p(doc, "RQ1 - Como pode a IA apoiar a governação da segurança da informação sob responsabilidade do CISO? A investigação mostra que a IA apoia a governação quando é integrada numa arquitetura de dados rastreável e controlada. No Virtual CISO, o assistente sintetiza informação, mas a governação assenta em entidades estruturadas, evidência e decisões registadas.")
    p(doc, "RQ2 - De que formas pode a IA auxiliar os CISOs na identificação, avaliação e mitigação de riscos? A resposta materializa-se na priorização contextual. O sistema combina EPSS, CISA KEV, criticidade do ativo, exposição, mitigação e evidência para ordenar vulnerabilidades de forma mais útil do que uma ordenação por CVSS. A explicação por contribuições permite justificar a decisão.")
    p(doc, "RQ3 - De que modo pode a IA apoiar o CISO na gestão da conformidade com normas e regulamentos? A resposta está na combinação de RAG rastreável, propagação de conformidade e drift. O sistema consegue recuperar obrigações normativas, relacioná-las com controlos internos, evidências e lacunas, e detetar regressões face a snapshots anteriores.")
    p(doc, "RQ4 - Que técnicas de IA são mais adequadas para tarefas de GRC? A resposta é híbrida. Regras determinísticas e modelos ponderados são adequados a decisões que exigem reprodutibilidade; RAG e LLM são adequados a síntese e consulta textual; ML preditivo é útil quando existe histórico suficiente. O Virtual CISO reflete esta distinção ao manter o modelo ponderado oficial, usar EPSS como sinal externo e condicionar o XGBoost+SHAP interno por prontidão.")

    h(doc, 2, "6.3. Principais Contributos")
    bullet(doc, "Um modelo integrado de trabalho para o CISO, operacionalizado por etapas, completude, bloqueios e próxima-melhor-ação.")
    bullet(doc, "Uma abordagem de priorização de vulnerabilidades sensível ao contexto organizacional, com explicação por fator.")
    bullet(doc, "Uma camada de rastreabilidade L4 que liga controlos externos, controlos internos, mecanismos e evidências.")
    bullet(doc, "Um RAG normativo e organizacional com fontes clicáveis ao nível de documento, artigo, parágrafo e página.")
    bullet(doc, "Um mecanismo de conformidade contínua baseado em snapshots e deteção de drift.")
    bullet(doc, "Um desenho governado para evolução de modelo ponderado para XGBoost+SHAP, evitando ativação prematura sem dados temporais.")

    h(doc, 2, "6.4. Limitações")
    p(doc, "A principal limitação é a ausência de validação longitudinal em organização real. Embora a aplicação esteja funcional e os casos de uso sejam demonstráveis, os resultados ainda não medem impacto operacional com CISOs em produção. Outra limitação é a não ativação do estimador interno XGBoost+SHAP: a implementação existe e as dependências estão disponíveis, mas o portão de prontidão ainda não permite ativação por falta de histórico suficiente. Por fim, os limiares de completude do ciclo do CISO dependem de juízo de desenho e devem ser parametrizados de acordo com apetite de risco.")

    h(doc, 2, "6.5. Trabalho Futuro")
    numbered(doc, "Executar validação com utilizadores CISO ou equipas de segurança, recolhendo métricas de utilidade, confiança e tempo de decisão.")
    numbered(doc, "Atingir o limiar de dados temporais e treinar o estimador interno XGBoost+SHAP, avaliando AUC, Brier e divergência face ao modelo ponderado em shadow mode.")
    numbered(doc, "Parametrizar os limiares do programa do CISO por apetite de risco e perfil regulatório.")
    numbered(doc, "Expandir o corpus RAG e a cobertura de fontes normativas, mantendo citações por parágrafo.")
    numbered(doc, "Integrar agendamentos automáticos de scans Nmap/Wazuh e snapshots de conformidade para reforçar a monitorização contínua.")
    numbered(doc, "Produzir relatórios executivos automáticos a partir do estado do programa, decisões registadas e drift detetado.")

    h(doc, 2, "6.6. Considerações Finais")
    p(doc, "O contributo central desta dissertação é demonstrar que uma plataforma Virtual CISO só é verdadeiramente útil quando deixa de ser um repositório de páginas e passa a orientar trabalho. A integração de dados é necessária, mas insuficiente; o valor emerge quando a aplicação responde ao CISO com uma sequência: onde está o bloqueio, qual é a próxima ação e por que razão essa ação é prioritária. O artefacto desenvolvido demonstra esta possibilidade de forma implementada, rastreável e alinhada com as lacunas identificadas na literatura.")


def references(doc: Document) -> None:
    h(doc, 1, "Referências citadas nos capítulos revistos")
    refs = [
        "Chen, T., & Guestrin, C. (2016). XGBoost: A scalable tree boosting system. Proceedings of the 22nd ACM SIGKDD International Conference on Knowledge Discovery and Data Mining.",
        "CISA. (2026). Known Exploited Vulnerabilities Catalog. Cybersecurity and Infrastructure Security Agency.",
        "FIRST. (2023). Exploit Prediction Scoring System (EPSS). Forum of Incident Response and Security Teams.",
        "Hevner, A. R., March, S. T., Park, J., & Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75-105.",
        "ISO/IEC. (2022). ISO/IEC 27001:2022 Information security, cybersecurity and privacy protection - Information security management systems.",
        "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. NeurIPS.",
        "Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems.",
        "NIST. (2024). The NIST Cybersecurity Framework (CSF) 2.0. National Institute of Standards and Technology.",
        "Parlamento Europeu e Conselho. (2022). Diretiva (UE) 2022/2555 relativa a medidas destinadas a garantir um elevado nível comum de cibersegurança na União (NIS2). Jornal Oficial da União Europeia.",
        "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A Design Science Research Methodology for Information Systems Research. Journal of Management Information Systems, 24(3), 45-77.",
    ]
    for ref in refs:
        p(doc, ref)


def annexes(doc: Document) -> None:
    h(doc, 1, "Anexo C - Demonstrações de código")
    p(doc, "Este anexo reúne excertos de código que sustentam as capacidades descritas nos capítulos 4 a 6. Cada secção indica a localização do ficheiro e as linhas aproximadas no repositório. O objetivo não é reproduzir todo o código, mas tornar auditável a ligação entre dissertação e implementação.")
    code_snippet(
        doc,
        "C.1. Fio condutor do CISO e etapas do programa",
        "backend/governance/services/program_service.py",
        185,
        225,
        "O serviço agrega sinais de ativos, risco, maturidade e drift, enriquece itens de trabalho com etapa e calcula foco atual e próxima ação.",
    )
    code_snippet(
        doc,
        "C.2. Sinal de ativos e ações da etapa Conhecer",
        "backend/governance/services/program_service.py",
        257,
        289,
        "A etapa Conhecer é calculada a partir de ativos oficiais, ativos em onboarding, tipo, responsável e classificação validada.",
    )
    code_snippet(
        doc,
        "C.3. Regra de onboarding e inventário fiável",
        "frontend/src/pages/assets/AssetOnboarding.tsx",
        217,
        235,
        "A função de frontend mantém no onboarding ativos sem responsável, sem tipo ou sem classificação validada.",
    )
    code_snippet(
        doc,
        "C.4. Inventário oficial apenas com ativos finalizados",
        "frontend/src/pages/assets/Inventory.tsx",
        136,
        142,
        "A página de inventário filtra ativos para mostrar apenas os que têm tipo e classificação validada, reforçando a distinção entre descoberta e inventário governável.",
    )
    code_snippet(
        doc,
        "C.5. Pesos do modelo de priorização contextual",
        "backend/risk/services/prioritization.py",
        17,
        33,
        "Os pesos incluem CVSS, EPSS, KEV e fatores organizacionais, refletindo a priorização multidimensional.",
    )
    code_snippet(
        doc,
        "C.6. Estimadores ponderado e XGBoost+SHAP",
        "backend/risk/services/prioritization.py",
        157,
        227,
        "A interface de estimador preserva o contrato score + decomposição, permitindo trocar o combinador sem alterar o frontend.",
    )
    code_snippet(
        doc,
        "C.7. Feature store da priorização",
        "backend/risk/services/prioritization.py",
        778,
        804,
        "Cada execução pode registar snapshots de features e desfechos para treino futuro do modelo interno.",
    )
    code_snippet(
        doc,
        "C.8. Portão de prontidão XGBoost+SHAP",
        "backend/risk/models/priority_model.py",
        65,
        150,
        "O modelo interno só pode ser ativado quando existem dados, ciclos, artefacto, métricas e dependências suficientes.",
    )
    code_snippet(
        doc,
        "C.9. Sincronização CISA KEV",
        "backend/risk/services/intel_service.py",
        132,
        185,
        "O serviço consulta o feed CISA KEV e marca CVEs existentes sem criar vulnerabilidades artificiais.",
    )
    code_snippet(
        doc,
        "C.10. Página de integração CISA KEV",
        "frontend/src/pages/admin/integrations/KEV.tsx",
        17,
        32,
        "A integração CISA KEV foi promovida para página própria de administração.",
    )
    code_snippet(
        doc,
        "C.11. Ingestão normativa NIS2 com citação por parágrafo",
        "backend/ciso_assistant/management/commands/ingest_nis2_normative.py",
        306,
        361,
        "A ingestão gera source_ref e metadados de artigo, parágrafo, página e ficheiro, permitindo citações clicáveis.",
    )
    code_snippet(
        doc,
        "C.12. Resolução de fonte RAG clicável",
        "backend/ciso_assistant/views.py",
        347,
        395,
        "O endpoint knowledge-source devolve a fonte original com metadados para visualização auditável.",
    )
    code_snippet(
        doc,
        "C.13. Serviço de rastreabilidade por framework",
        "backend/governance/services/traceability_service.py",
        400,
        449,
        "A rastreabilidade por framework agrega secções, controlos, mapeamentos, scores, gaps e evidência.",
    )
    code_snippet(
        doc,
        "C.14. Deteção determinística de drift",
        "backend/governance/services/security_posture_drift.py",
        115,
        156,
        "O overview de drift combina regressões de controlos, regressões de exposição e vulnerabilidades novas com severidade e recomendações.",
    )

    h(doc, 1, "Anexo D - Registo de capturas de ecrã")
    p(doc, "As imagens foram recolhidas da aplicação local em execução, numa sessão autenticada, e guardadas em docs/imagens/.")
    table(doc, ["Ficheiro", "Ecrã", "Uso no texto"], [
        ["fig-4-01-ciclo-programa-ciso.png", "Mission Control / modo programático", "Fio condutor do CISO e etapa atual."],
        ["fig-4-02-onboarding-ativos.png", "Entrada de ativos", "Regra de inventário fiável e backlog de onboarding."],
        ["fig-5-01-priorizacao-contextual.png", "Priorização contextual", "UC1 - CVSS idêntico, prioridade diferente."],
        ["fig-5-02-modelo-priorizacao.png", "Modelo de Priorização (IA)", "Portão XGBoost+SHAP e feature store."],
        ["fig-5-03-cisa-kev-integracao.png", "Integração CISA KEV", "Fonte externa de exploração conhecida."],
        ["fig-5-04-rastreabilidade-framework.png", "Explorador de rastreabilidade", "Cadeia L4 por framework."],
        ["fig-5-05-fonte-nis2-rag.png", "Fonte normativa RAG", "UC2 - citação NIS2 por artigo/parágrafo."],
        ["fig-5-06-drift-conformidade.png", "Drift de conformidade", "UC3 - regressão de postura."],
    ], widths=[2.4, 2.0, 2.1])


def main() -> None:
    doc = Document()
    style_document(doc)
    add_cover(doc)
    chapter_4(doc)
    doc.add_page_break()
    chapter_5(doc)
    doc.add_page_break()
    chapter_6(doc)
    doc.add_page_break()
    references(doc)
    doc.add_page_break()
    annexes(doc)
    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    main()
