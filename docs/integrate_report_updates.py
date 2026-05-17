from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\virtual_ciso\virtual_CISO_V1_2\docs")
SOURCE_PATH = ROOT / "Relatorio_Rui_Agostinho_15_05_2026.docx"
OUTPUT_PATH = ROOT / "Relatorio_Rui_Agostinho_17_05_2026.docx"


def find_heading_index(doc: Document, heading_text: str) -> int:
    for idx, paragraph in enumerate(doc.paragraphs):
        if paragraph.text.strip() == heading_text:
            return idx
    raise ValueError(f"Heading not found: {heading_text}")


def insert_paragraph_before(paragraph: Paragraph, text: str, style: str) -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addprevious(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if text:
        new_paragraph.add_run(text)
    if style:
        new_paragraph.style = style
    return new_paragraph


def set_paragraph(paragraph: Paragraph, text: str, style: str) -> None:
    paragraph.text = text
    if style:
        paragraph.style = style


def replace_block(doc: Document, start_heading: str, end_heading: str, items: list[tuple[str, str]]) -> None:
    start_idx = find_heading_index(doc, start_heading)
    end_idx = find_heading_index(doc, end_heading)
    existing = list(doc.paragraphs[start_idx + 1 : end_idx])

    for paragraph, (style, text) in zip(existing, items):
        set_paragraph(paragraph, text, style)

    if len(existing) > len(items):
        for paragraph in existing[len(items) :]:
            set_paragraph(paragraph, "", paragraph.style.name if paragraph.style else "Normal")

    if len(items) > len(existing):
        anchor = doc.paragraphs[find_heading_index(doc, end_heading)]
        for style, text in items[len(existing) :]:
            insert_paragraph_before(anchor, text, style)


CHAPTER_UPDATES: list[tuple[str, str, list[tuple[str, str]]]] = [
    (
        "IA Generativa e RAG Híbrido",
        "Orquestrador Principal",
        [
            (
                "Normal",
                "A componente de IA generativa do Virtual CISO foi concebida como um mecanismo de apoio à interpretação e à decisão, e não como um substituto da lógica determinística da plataforma. A sua função consiste em transformar dados técnicos, normativos, organizacionais e evidenciais em respostas compreensíveis, recomendações contextualizadas e justificações adequadas a diferentes perfis de utilizador. Esta componente encontra-se implementada na aplicação Django ciso_assistant (backend/ciso_assistant/) e responde particularmente às necessidades identificadas na revisão da literatura quanto à explicação, síntese e contextualização da informação de GRC.",
            ),
            (
                "Normal",
                "Contudo, a utilização isolada de modelos de linguagem em contextos de cibersegurança e governação apresenta limitações conhecidas: alucinação factual, fragilidade em consultas numéricas ou agregadas, ausência de conhecimento sobre o estado real da organização e dificuldade em sustentar respostas auditáveis. Num domínio em que o utilizador pode necessitar de justificar decisões perante auditorias, órgãos de gestão ou equipas técnicas, tais limitações tornariam inadequada uma arquitetura puramente generativa.",
            ),
            (
                "Normal",
                "Para mitigar estes riscos, o Virtual CISO adota uma arquitetura de RAG híbrido. Em vez de delegar toda a resposta ao LLM, o sistema articula três fontes de inteligência complementares: consultas estruturadas à base de dados relacional, recuperação semântica de conhecimento documental indexado e geração final em linguagem natural. Deste modo, a componente generativa opera condicionada por contexto efetivamente existente no sistema, reforçando a relevância, a rastreabilidade e a utilidade prática das respostas produzidas.",
            ),
        ],
    ),
    (
        "Orquestrador Principal",
        "Routing de Modelos (LLMRouter)",
        [
            (
                "Normal",
                "A orquestração do pipeline é realizada pela classe QueryOrchestrator (backend/ciso_assistant/services/orchestrator.py), através do método process_query(query, history, filters). Esta classe coordena o fluxo completo da pergunta, desde a classificação inicial até à construção da resposta final, funcionando como ponto de articulação entre os componentes determinísticos, o mecanismo RAG e o serviço LLM.",
            ),
            ("List Paragraph", "LLMRouter - classifica a intenção da pergunta, estima a necessidade de RAG e seleciona o perfil de resposta adequado;"),
            ("List Paragraph", "StructuredQueryService - responde diretamente a consultas factuais, contagens e agregações executadas via Django ORM;"),
            ("List Paragraph", "VulnerabilityPrioritizationService - obtém rankings determinísticos de vulnerabilidades com base na engine de priorização descrita na Secção 4.6.2;"),
            ("List Paragraph", "SemanticRetrievalService - recupera segmentos documentais semanticamente relevantes a partir do índice vetorial;"),
            ("List Paragraph", "PromptBuilder e OllamaClient - constroem o contexto final e executam a geração em linguagem natural no modelo selecionado."),
            (
                "Normal",
                "A decisão arquitetural mais relevante reside no facto de o orquestrador privilegiar sempre a rota mais controlável antes de invocar geração livre. Perguntas estritamente estruturadas são respondidas sem recurso ao LLM; perguntas de priorização usam primeiro a engine formal de scoring; e, quando a consulta requer contexto interno mas esse contexto não é recuperado com qualidade suficiente, o sistema devolve uma resposta explícita de insuficiência de dados em vez de inferir conteúdo não suportado. Adicionalmente, no fluxo de priorização, foi introduzido um fallback textual determinístico para os casos em que o serviço LLM falha ou excede o tempo de resposta configurado.",
            ),
        ],
    ),
    (
        "Routing de Modelos (LLMRouter)",
        "Recuperação Semântica e Embeddings",
        [
            (
                "Normal",
                "A seleção do comportamento do assistente é realizada pela classe LLMRouter (backend/ciso_assistant/services/llm_router.py). O router não escolhe apenas um modelo: decide também o tipo de tarefa, se a pergunta deve ativar recuperação RAG e qual o perfil de geração mais adequado ao contexto. Esta função é central para equilibrar qualidade, latência e controlo, sobretudo num ambiente em que coexistem perguntas executivas, consultas factuais, pedidos de priorização e questões conceptuais.",
            ),
            (
                "Normal",
                "A implementação segue uma estratégia híbrida de decisão. Em primeiro lugar, aplica regras estritas para consultas estruturadas e para casos de precedência elevada, como priorização de vulnerabilidades ou perguntas conceptuais gerais. Em segundo lugar, utiliza um motor de pontuação baseado em padrões linguísticos para os tipos de tarefa mais recorrentes. Só em situações não resolvidas por estas camadas é usado um classificador LLM para produzir uma decisão estruturada. Esta composição permite combinar previsibilidade e auditabilidade com flexibilidade semântica.",
            ),
            ("Caption", "Tabela 21 - Mapeamento de tarefas a modelos no LLMRouter"),
            (
                "Normal",
                "Do ponto de vista arquitetural, esta abordagem é particularmente adequada ao domínio de GRC porque evita tratar todas as perguntas como problemas generativos equivalentes. Ao distinguir explicitamente tarefas determinísticas, perguntas de priorização e questões que beneficiam de síntese em linguagem natural, o router reduz o risco de alucinação, melhora o tempo de resposta em cenários estruturados e cria uma base mensurável para a avaliação apresentada no Capítulo 5.",
            ),
        ],
    ),
    (
        "Recuperação Semântica e Embeddings",
        "Sincronização e Consistência do Índice Semântico",
        [
            (
                "Normal",
                "A recuperação semântica é implementada pela classe SemanticRetrievalService (backend/ciso_assistant/services/semantic_retrieval.py), que executa pesquisa por similaridade vetorial sobre segmentos de conhecimento armazenados em PostgreSQL com a extensão pgvector. Esta camada é usada quando a pergunta exige enquadramento documental, políticas internas, controlos, procedimentos ou explicações cuja formulação não se encontra diretamente codificada na base relacional.",
            ),
            (
                "Normal",
                "Os embeddings são gerados pelo serviço de embeddings configurado na plataforma e associados às instâncias de KnowledgeChunk. Durante a pesquisa, o sistema calcula o embedding da pergunta, filtra segmentos semanticamente pesquisáveis e devolve o conjunto Top-K mais relevante. Na implementação atual, os segmentos sem embedding válido são explicitamente excluídos do retrieval, evitando ruído e resultados degradados por indexação incompleta.",
            ),
            (
                "Normal",
                "O resultado desta pesquisa não é usado isoladamente. Os segmentos recuperados são normalizados, combinados com dados estruturados e integrados num contexto único pelo PromptBuilder. Este desenho permite que o assistente trabalhe com conhecimento documental contextualizado pela realidade operacional da organização, em vez de responder apenas com base em texto normativo genérico.",
            ),
        ],
    ),
    (
        "Sincronização e Consistência do Índice Semântico",
        "Pipeline Completo",
        [
            (
                "Normal",
                "A utilidade de uma arquitetura RAG depende não apenas da qualidade dos embeddings, mas também da consistência entre a base operacional e o índice semântico. Se uma política, um mecanismo, uma evidência ou uma vulnerabilidade forem alterados na aplicação sem atualização correspondente do seu registo textual pesquisável, o sistema corre o risco de recuperar conhecimento desatualizado e produzir respostas desalinhadas com o estado real da organização.",
            ),
            (
                "Normal",
                "Para reduzir esse risco, o Virtual CISO combina duas modalidades de atualização do conhecimento indexado. A primeira corresponde à sincronização incremental orientada a eventos, adequada a entidades curadas manualmente e com impacto direto na governação e conformidade, como controlos, mecanismos, políticas, procedimentos e findings. Sempre que estes elementos são criados ou alterados, o respetivo conteúdo semântico pode ser atualizado após persistência bem-sucedida na base de dados.",
            ),
            (
                "Normal",
                "A segunda modalidade corresponde à indexação em lote, utilizada para reconstruções controladas do índice ou para entidades potencialmente volumosas, como ativos descobertos automaticamente e vulnerabilidades importadas de fontes externas. Esta separação evita introduzir latência excessiva nas operações de ingestão e permite tratar cargas massivas de forma administrável.",
            ),
            (
                "Normal",
                "A existência de mecanismos explícitos de reindexação reforça a governabilidade da componente RAG. A atualização do conhecimento pesquisável deixa de ser uma consequência implícita e opaca do sistema para passar a constituir um processo controlado, observável e repetível, coerente com o tipo de exigência de auditabilidade esperado em contexto GRC.",
            ),
            (
                "Normal",
                "No estado atual do protótipo, parte desta sincronização continua a ser realizada de forma síncrona, o que é aceitável para entidades de baixo volume e edição manual. Em ambiente produtivo, contudo, a evolução natural passará pelo desacoplamento do processo de geração de embeddings através de filas ou workers assíncronos, reduzindo o impacto na experiência de edição e aumentando a resiliência perante indisponibilidade temporária do serviço de embeddings.",
            ),
            (
                "Normal",
                "Esta decisão de separar atualização incremental e reindexação em lote é relevante também para a qualidade do assistente. Ao explicitar quando e como o conhecimento passa a estar semanticamente disponível para pesquisa, o sistema reduz o risco de respostas baseadas em contexto parcial e melhora a previsibilidade da camada conversacional.",
            ),
            (
                "Normal",
                "Assim, a consistência do índice semântico é tratada como requisito arquitetural do assistente e não como detalhe meramente técnico. Essa opção contribui diretamente para a robustez das respostas, para a sua rastreabilidade e para a confiança do utilizador na informação devolvida.",
            ),
        ],
    ),
    (
        "Pipeline Completo",
        "Rastreabilidade das Respostas",
        [
            (
                "Normal",
                "O pipeline completo da componente de IA híbrida pode ser entendido como uma sequência de cinco etapas encadeadas, ilustradas na Figura 17 e materializadas diretamente em QueryOrchestrator.process_query().",
            ),
            ("Caption", "Figura 17 - Arquitetura de RAG híbrido do Virtual CISO"),
            (
                "List Paragraph",
                "Receção e classificação da pergunta - o LLMRouter identifica a intenção dominante, estima a necessidade de recuperação RAG e escolhe a rota de execução mais adequada;",
            ),
            (
                "List Paragraph",
                "Recuperação estruturada - StructuredQueryService ou VulnerabilityPrioritizationService consultam a base relacional para responder a perguntas factuais, agregadas ou de ranking;",
            ),
            (
                "List Paragraph",
                "Recuperação semântica - SemanticRetrievalService pesquisa segmentos documentais relevantes no índice vetorial com base na similaridade do significado;",
            ),
            (
                "List Paragraph",
                "Construção do contexto - PromptBuilder agrega dados estruturados, fontes semânticas, histórico de conversa e instruções de resposta num único contexto coerente;",
            ),
            (
                "List Paragraph",
                "Geração final - OllamaClient invoca o modelo selecionado, com perfil de geração ajustado ao tipo de tarefa, produzindo uma resposta em linguagem natural com foco em utilidade e fundamentação.",
            ),
            (
                "Normal",
                "A relevância desta cadeia não reside apenas no uso de RAG, mas na forma como o contexto é montado. A informação sobre mecanismos de implementação, evidência existente, findings, ativos e controlos aplicáveis permite que a resposta final deixe de ser uma recomendação abstrata e passe a refletir a realidade operacional conhecida da organização.",
            ),
        ],
    ),
    (
        "Rastreabilidade das Respostas",
        "Interface e Apresentação de Resultados",
        [
            (
                "Normal",
                "A rastreabilidade das respostas é assegurada pela classe SourceNormalizer (backend/ciso_assistant/services/source_normalizer.py), responsável por transformar os elementos usados na resposta numa estrutura uniforme de fontes. Cada fonte inclui, sempre que aplicável, título, tipo de origem, entidade relacionada, conteúdo utilizado e peso de relevância, permitindo que a interface apresente ao utilizador não apenas a resposta, mas também os seus suportes.",
            ),
            (
                "Normal",
                "Esta capacidade é decisiva para distinguir uma recomendação fundamentada de uma resposta meramente plausível. No Virtual CISO, a resposta conversacional pode ser associada a vulnerabilidades, ativos, controlos, mecanismos, evidência, findings ou excertos documentais concretos. O utilizador consegue, assim, validar a recomendação, seguir a cadeia explicativa subjacente e exercer controlo humano efetivo sobre a decisão suportada pelo assistente.",
            ),
        ],
    ),
    (
        "Critérios e Método de Avaliação",
        "Avaliação dos Requisitos Funcionais",
        [
            (
                "Normal",
                "A avaliação do Virtual CISO foi conduzida em coerência com a lógica de Design Science Research, privilegiando a análise da utilidade, da qualidade e da adequação do artefacto enquanto prova de conceito. Não se tratou apenas de verificar se a aplicação funciona tecnicamente, mas de apreciar se a solução responde de forma plausível ao problema de investigação identificado: a necessidade de apoiar o CISO com uma visão integrada, contextualizada e rastreável sobre risco, conformidade, mecanismos e evidência.",
            ),
            (
                "Normal",
                "Neste sentido, a avaliação não incidiu exclusivamente sobre o assistente de IA. Incidiu sobre o comportamento integrado do artefacto, incluindo modelação do contexto organizacional, inventário de ativos, correlação entre vulnerabilidades e controlos, scoring de risco residual, avaliação multi-framework, navegação entre níveis de abstração e explicabilidade das recomendações devolvidas ao utilizador.",
            ),
            (
                "Normal",
                "Foram mobilizados quatro instrumentos complementares. O primeiro correspondeu à demonstração funcional do artefacto, organizada em casos de utilização representativos. O segundo consistiu na análise de cobertura dos requisitos funcionais e não funcionais definidos no Capítulo 4. O terceiro correspondeu a uma avaliação específica do LLMRouter, com dataset rotulado e métricas quantitativas. O quarto incidiu sobre o pipeline completo do assistente, testando routing, bypass estruturado, recuperação RAG, geração de resposta e fontes devolvidas.",
            ),
            (
                "Normal",
                "A combinação destes instrumentos permitiu articular evidência qualitativa e quantitativa. Assim, a avaliação pôde captar simultaneamente a coerência arquitetural da solução, a sua utilidade prática nos fluxos principais e o comportamento específico da componente de IA híbrida, sem reduzir a apreciação do Virtual CISO à performance isolada de um modelo de linguagem.",
            ),
        ],
    ),
    (
        "Avaliação dos Requisitos Funcionais",
        "Avaliação dos Requisitos Não Funcionais",
        [
            (
                "Normal",
                "A análise de cobertura dos requisitos funcionais é sintetizada na Tabela 29. Para cada requisito, foi considerado o grau em que a capacidade correspondente se encontra efetivamente demonstrada no cenário de avaliação e na implementação do protótipo. A leitura da tabela deve ser entendida em articulação com os casos de utilização descritos na secção 5.3, uma vez que a utilidade do artefacto emerge sobretudo da integração entre módulos e não de funcionalidades isoladas.",
            ),
            ("Normal", "Tabela 29 - Avaliação da cobertura dos requisitos funcionais"),
            (
                "Normal",
                "No conjunto do cenário analisado, o Virtual CISO demonstra cobertura forte das capacidades centrais do artefacto: modelação do contexto organizacional, gestão de ativos, associação entre vulnerabilidades e ativos, avaliação de conformidade multi-framework, cálculo de risco residual, representação explícita de mecanismos de implementação e apoio à priorização. O requisito associado à geração de recomendações por IA mantém natureza evolutiva, não por ausência de implementação, mas porque a sua qualidade final depende de refinamento contínuo do conhecimento disponível, dos prompts e dos perfis de geração utilizados.",
            ),
            (
                "Normal",
                "A avaliação funcional do subsistema conversacional reforça esta leitura. O LLMRouter, após afinamento das regras de precedência, atingiu 100,00% de acerto em task_type, 100,00% em needs_rag e 100,00% de acerto exato no dataset avaliado. Já o assistente, quando testado ponta a ponta, apresentou 100,00% de correção funcional nos cenários estruturados e nos cenários de avaliação usados para validação do fluxo conversacional. Estes resultados não substituem validação com utilizadores, mas demonstram que a componente de IA está operacionalmente integrada nas funções do artefacto.",
            ),
        ],
    ),
    (
        "Avaliação dos Requisitos Não Funcionais",
        "Análise Comparativa com Soluções Existentes",
        [
            (
                "Normal",
                "A avaliação dos requisitos não funcionais centra-se nas propriedades transversais que sustentam a confiança e a governabilidade do artefacto. A Tabela 30 sintetiza a forma como estes requisitos são operacionalizados, com especial destaque para explicabilidade, rastreabilidade, controlo humano, segurança da informação e capacidade de evolução arquitetural.",
            ),
            ("Normal", "Tabela 30 - Avaliação da cobertura dos requisitos não funcionais"),
            (
                "Normal",
                "Os requisitos de explicabilidade e rastreabilidade encontram-se particularmente bem suportados porque foram incorporados como decisões estruturais do sistema e não como adições tardias. A navegação entre ativos, vulnerabilidades, controlos, mecanismos, evidência e findings; a normalização das fontes do assistente; e a distinção entre respostas determinísticas, priorização formal e geração contextualizada constituem mecanismos concretos de controlo e justificação. De modo semelhante, o controlo humano é reforçado pelo facto de o sistema poder devolver insuficiência de contexto ou recorrer a bypasses estruturados em vez de fabricar respostas plausíveis mas não sustentadas.",
            ),
            (
                "Normal",
                "Em termos de desempenho, a avaliação mostrou um comportamento assimétrico entre fluxos estruturados e generativos. As respostas baseadas em bypass ORM apresentaram latências muito baixas, na ordem de dezenas a poucas centenas de milissegundos. Nos cenários generativos, observou-se inicialmente um gargalo associado ao serviço LLM remoto. Após o aumento controlado do timeout, a introdução de fallback determinístico para priorização e o ajuste do perfil de geração por tarefa, a pergunta 'O que faz um CISO?' passou de aproximadamente 96 segundos para cerca de 22 segundos, e a pergunta 'Prioriza as vulnerabilidades mais urgentes' desceu de cerca de 174 segundos para aproximadamente 86 segundos. Estes resultados evidenciam melhoria relevante, mas mostram também que escalabilidade e desempenho continuam a constituir dimensão a aprofundar em trabalho futuro.",
            ),
        ],
    ),
    (
        "Discussão dos Resultados",
        "Limitações da Avaliação",
        [
            (
                "Normal",
                "Os resultados observados permitem retirar quatro conclusões principais sobre o Virtual CISO enquanto artefacto de apoio à decisão. A primeira é a viabilidade arquitetural da proposta. A integração entre dados técnicos, requisitos normativos, mecanismos de implementação, evidência e capacidades conversacionais mostrou-se exequível num único artefacto coerente, capaz de suportar leitura executiva e análise operacional.",
            ),
            (
                "Normal",
                "A segunda conclusão é a utilidade do modelo de pontuação multidimensional para a priorização. A demonstração confirmou que vulnerabilidades com severidade técnica idêntica podem assumir prioridades distintas quando são contextualizadas por criticidade do ativo, exposição, estado dos mecanismos, evidência e relevância regulatória. Esta capacidade aproxima a plataforma de uma lógica de apoio à decisão e afasta-a de dashboards meramente descritivos.",
            ),
            (
                "Normal",
                "A terceira conclusão diz respeito à componente de IA híbrida. A avaliação do LLMRouter e do assistente indica que a combinação entre regras, ORM, RAG e LLM é mais adequada ao domínio estudado do que uma abordagem puramente generativa. O router passou de 90,48% para 100,00% de acerto exato no dataset avaliado após refinamento das regras, e o assistente revelou melhoria mensurável de comportamento e de latência após ajustes de timeout, fallback e perfis de geração.",
            ),
            (
                "Normal",
                "A quarta conclusão é que a robustez da experiência conversacional depende menos de um modelo isolado do que da qualidade da integração entre dados estruturados, conhecimento indexado, regras de orquestração e controlo humano. Em termos práticos, o valor do assistente emerge quando este é tratado como extensão de um sistema de GRC e não como interface genérica desacoplada da realidade operacional da organização.",
            ),
            (
                "Normal",
                "Estas conclusões devem, contudo, ser interpretadas no quadro de uma prova de conceito. A avaliação sustenta a adequação conceptual e a viabilidade técnica do artefacto, mas não constitui ainda evidência de impacto organizacional ou de adoção em ambiente real. Ainda assim, oferece uma base suficientemente sólida para defender o Virtual CISO como contributo relevante no espaço entre plataformas de GRC, apoio à decisão e IA aplicada.",
            ),
        ],
    ),
    (
        "Limitações da Avaliação",
        "Ameaças à Validade",
        [
            (
                "Normal",
                "A avaliação realizada apresenta limitações que importa explicitar para enquadrar adequadamente os resultados. A primeira decorre da natureza do cenário de demonstração. Embora o caso tenha sido construído para ser plausível e coerente com uma entidade essencial, não corresponde a uma organização real observada longitudinalmente, o que limita a representatividade externa dos resultados.",
            ),
            (
                "Normal",
                "A segunda limitação prende-se com a ausência de avaliação direta por utilizadores finais. A demonstração e os testes foram conduzidos pelo investigador, sem participação sistemática de CISOs, equipas de risco, responsáveis de conformidade ou outros perfis-alvo. Consequentemente, continuam em aberto dimensões como perceção de utilidade, confiança, usabilidade e integração no trabalho quotidiano.",
            ),
            (
                "Normal",
                "A terceira limitação respeita à própria componente de IA. Apesar dos resultados positivos obtidos no router e nos cenários funcionais do assistente, os datasets utilizados são ainda reduzidos e orientados a cenários representativos da prova de conceito. Isso significa que os valores quantitativos obtidos devem ser lidos como evidência inicial de robustez e não como validação estatística definitiva.",
            ),
            (
                "Normal",
                "A quarta limitação relaciona-se com desempenho e infraestrutura. A avaliação evidenciou melhoria relevante após otimizações, mas também mostrou que a qualidade percebida do assistente depende da disponibilidade do serviço LLM, da latência de geração e da maturidade do índice semântico. Em contexto produtivo, estes fatores exigiriam observação continuada, testes de carga e monitorização operacional.",
            ),
            (
                "Normal",
                "A quinta limitação é transversal a todo o artefacto: a dependência da qualidade e da completude dos dados de entrada. Como qualquer sistema de apoio à decisão, o Virtual CISO só pode produzir análises tão fiáveis quanto os dados sobre ativos, vulnerabilidades, controlos, mecanismos e evidência que lhe são fornecidos. Em organizações com baixa maturidade de governação destes dados, o impacto do artefacto tenderá a ser inferior ao observado no cenário de demonstração.",
            ),
        ],
    ),
    (
        "Ameaças à Validade",
        "Síntese da Demonstração e Avaliação",
        [
            (
                "Normal",
                "Para além das limitações operacionais, importa considerar as principais ameaças à validade dos resultados. Em termos de validade de construção, existe o risco de algumas entidades e relações implementadas no artefacto representarem apenas parcialmente os conceitos teóricos que procuram operacionalizar. A própria noção de mecanismo de implementação, central nesta dissertação, poderá assumir configurações distintas consoante o contexto organizacional, o setor ou o referencial normativo adotado.",
            ),
            (
                "Normal",
                "No plano da validade interna, a principal ameaça resulta do facto de a demonstração e parte relevante da avaliação terem sido conduzidas pelo próprio investigador, o que pode introduzir viés na escolha dos cenários, na formulação das queries e na interpretação dos resultados. Esta ameaça é mitigada parcialmente pela explicitação prévia dos requisitos, pela separação entre avaliação funcional do artefacto e avaliação específica do router/assistente, e pela utilização de comandos reprodutíveis para os testes quantitativos.",
            ),
            (
                "Normal",
                "Quanto à validade externa, os resultados não podem ser generalizados automaticamente a outros contextos organizacionais, setores ou perfis de maturidade. A arquitetura modular e a modelação multi-framework sugerem potencial de adaptação, mas essa hipótese carece de validação empírica em ambientes distintos. Acresce que o comportamento da componente de IA poderá variar em função da infraestrutura LLM disponível, da cobertura da base de conhecimento e do volume de dados indexados.",
            ),
            (
                "Normal",
                "Estas ameaças não anulam a utilidade dos resultados obtidos, mas delimitam o seu alcance interpretativo. A sua explicitação é importante porque reforça a transparência metodológica da investigação e clarifica em que medida a avaliação apresentada deve ser entendida como evidência de viabilidade e não como certificação de maturidade final do artefacto.",
            ),
        ],
    ),
    (
        "Síntese da Demonstração e Avaliação",
        "Conclusões e Trabalho Futuro",
        [
            (
                "Normal",
                "O presente capítulo apresentou a demonstração e avaliação do Virtual CISO num cenário representativo, em coerência com as fases finais da metodologia DSR adotada. A demonstração confirmou que o artefacto suporta, de forma integrada, modelação organizacional, gestão de ativos, correlação entre vulnerabilidades e ativos, avaliação de conformidade multi-framework, scoring de risco residual, priorização de ações e geração de respostas assistidas por IA com fundamento verificável.",
            ),
            (
                "Normal",
                "A avaliação mostrou que o valor do Virtual CISO não reside numa única funcionalidade, mas na articulação entre os seus módulos. Os testes funcionais e arquiteturais indicaram cobertura forte dos requisitos centrais definidos no Capítulo 4, enquanto a análise específica da componente de IA evidenciou evolução mensurável após refinamento: o LLMRouter atingiu 100,00% de acerto no dataset utilizado e o assistente melhorou substancialmente a latência e a robustez em cenários representativos.",
            ),
            (
                "Normal",
                "A discussão dos resultados confirmou três ideias-chave: a viabilidade da proposta enquanto plataforma integrada de apoio à decisão em GRC; a utilidade da representação explícita dos mecanismos de implementação e do modelo de pontuação multidimensional; e a adequação de uma arquitetura híbrida que combina lógica determinística, recuperação contextual e geração em linguagem natural. Em conjunto, estes elementos respondem de forma consistente ao problema de investigação que motivou o trabalho.",
            ),
            (
                "Normal",
                "Embora subsistam limitações relevantes — nomeadamente ausência de validação com utilizadores, datasets de avaliação ainda reduzidos e necessidade de aprofundar desempenho e escalabilidade — os resultados obtidos sustentam a conclusão de que o Virtual CISO constitui um contributo original e promissor para o cruzamento entre plataformas de GRC, apoio à decisão em cibersegurança e IA explicável aplicada a contexto organizacional.",
            ),
        ],
    ),
]


def main() -> None:
    shutil.copyfile(SOURCE_PATH, OUTPUT_PATH)
    doc = Document(OUTPUT_PATH)

    for start_heading, end_heading, items in CHAPTER_UPDATES:
        replace_block(doc, start_heading, end_heading, items)

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
