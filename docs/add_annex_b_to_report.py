from __future__ import annotations

import shutil
from pathlib import Path

from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph


ROOT = Path(r"D:\virtual_ciso\virtual_CISO_V1_2\docs")
SOURCE_PATH = ROOT / "Relatorio_Rui_Agostinho_17_05_2026.docx"
OUTPUT_PATH = ROOT / "Relatorio_Rui_Agostinho_18_05_2026.docx"


def find_heading(doc: Document, text: str) -> Paragraph:
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == text:
            return paragraph
    raise ValueError(f"Heading not found: {text}")


def insert_paragraph_after(paragraph: Paragraph, text: str, style: str = "Normal") -> Paragraph:
    new_p = OxmlElement("w:p")
    paragraph._p.addnext(new_p)
    new_paragraph = Paragraph(new_p, paragraph._parent)
    if style:
        new_paragraph.style = style
    if text:
        new_paragraph.add_run(text)
    return new_paragraph


def append_after(paragraph: Paragraph, additions: list[tuple[str, str]]) -> None:
    current = paragraph
    for style, text in additions:
        current = insert_paragraph_after(current, text, style)


ANNEX_ITEMS = [
    ("Heading 2", "B.1. Modelação de governação e mecanismos"),
    (
        "Normal",
        "Os ficheiros seguintes correspondem às referências explícitas do Capítulo 4 relativas à modelação de mecanismos, políticas e controlos normativos.",
    ),
    (
        "List Paragraph",
        "backend/governance/models/mechanism.py - define as entidades principais associadas ao catálogo e à instanciação de mecanismos de implementação.",
    ),
    (
        "List Paragraph",
        "backend/governance/models/policy_management.py - inclui a entidade ImplementationMechanism, usada para relacionar políticas internas com mecanismos implementados.",
    ),
    (
        "List Paragraph",
        "backend/governance/models/framework.py - modela referenciais e secções de framework utilizados na governação e na conformidade.",
    ),
    (
        "List Paragraph",
        "backend/governance/models/control.py - representa os controlos associados aos referenciais suportados pela plataforma.",
    ),
    (
        "List Paragraph",
        "backend/governance/services/conformidade_gap_engine.py - implementa a lógica de avaliação de conformidade agregada e por controlo.",
    ),
    ("Heading 2", "B.2. Risco, ativos e priorização"),
    (
        "Normal",
        "Esta secção agrega os ficheiros citados no texto para suportar a descrição do modelo de dados de risco, da correlação entre ativos e vulnerabilidades e das engines analíticas.",
    ),
    (
        "List Paragraph",
        "backend/risk/models/vulnerability.py - contém a entidade AssetVulnerability, usada para relacionar ativos, software e vulnerabilidades com estado contextual.",
    ),
    (
        "List Paragraph",
        "backend/risk/models/risk.py - define configurações e entidades de suporte ao cálculo de risco inerente e residual.",
    ),
    (
        "List Paragraph",
        "backend/risk/services/risk_engine.py - implementa a classe RiskEngineService e o cálculo de score de risco normalizado.",
    ),
    (
        "List Paragraph",
        "backend/risk/services/prioritization.py - implementa a VulnerabilityScoringEngine responsável pela priorização multidimensional de vulnerabilidades.",
    ),
    ("Heading 2", "B.3. Assistente de IA e arquitetura RAG híbrida"),
    (
        "Normal",
        "Os ficheiros seguintes sustentam a descrição do pipeline conversacional e da componente de IA híbrida. Incluem tanto os elementos referidos explicitamente no texto principal como os ficheiros instrumentais usados na afinação e avaliação da solução.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/orchestrator.py - coordena o fluxo completo da pergunta, incluindo routing, bypass estruturado, RAG, fallback e geração final.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/llm_router.py - implementa a classificação de task_type, a decisão de ativação de RAG e a seleção do perfil de resposta.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/semantic_retrieval.py - executa recuperação semântica sobre o índice vetorial e filtra segmentos pesquisáveis.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/source_normalizer.py - uniformiza as fontes devolvidas com cada resposta do assistente.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/ollama_client.py - encapsula o acesso ao serviço LLM, controlo de timeout e tratamento de falhas do runtime.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/prompt_builder.py - constrói o contexto textual final fornecido ao modelo de linguagem.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/services/retrieval/embedding_service.py - suporta a geração de embeddings usados no índice semântico.",
    ),
    (
        "List Paragraph",
        "backend/core/settings.py - centraliza parâmetros de configuração relevantes para a componente LLM, incluindo timeout e perfis de geração.",
    ),
    ("Heading 2", "B.4. Interface de utilizador"),
    (
        "Normal",
        "Os ficheiros seguintes suportam a descrição da camada de apresentação e da integração do assistente no frontend.",
    ),
    (
        "List Paragraph",
        "frontend/src/pages/MissionControl.tsx - implementa o dashboard executivo usado como ponto de entrada para a leitura agregada da postura de segurança.",
    ),
    (
        "List Paragraph",
        "frontend/src/components/chat/ChatWidget.tsx - integra o assistente conversacional na interface principal e apresenta respostas com fontes associadas.",
    ),
    (
        "List Paragraph",
        "frontend/src/pages/Assistant.tsx - suporta a página dedicada ao assistente e à experiência conversacional alargada.",
    ),
    ("Heading 2", "B.5. Artefactos de avaliação do router e do assistente"),
    (
        "Normal",
        "Para reforçar a reprodutibilidade da avaliação apresentada no Capítulo 5, são indicados também os comandos e datasets usados na validação quantitativa e funcional da componente de IA.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/management/commands/evaluate_llm_router.py - executa a avaliação quantitativa do LLMRouter sobre um conjunto rotulado de queries.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/management/commands/evaluate_assistant.py - avalia o pipeline completo do assistente, incluindo routing, retrieval, geração e fontes devolvidas.",
    ),
    (
        "List Paragraph",
        "backend/ciso_assistant/management/commands/diagnose_ai.py - fornece diagnóstico operacional da configuração do serviço de IA usado pela plataforma.",
    ),
    (
        "List Paragraph",
        "docs/llmrouter_test_queries.json - dataset de perguntas usado na avaliação específica do router.",
    ),
    (
        "List Paragraph",
        "docs/assistant_eval_queries.json - dataset de cenários usado na avaliação funcional do assistente.",
    ),
]


def main() -> None:
    shutil.copyfile(SOURCE_PATH, OUTPUT_PATH)
    doc = Document(OUTPUT_PATH)

    structure_heading = find_heading(doc, "Estrutura do Projeto")
    append_after(
        structure_heading,
        [
            (
                "Normal",
                "Para reforçar a verificabilidade técnica da dissertação, o Anexo B consolida as referências explícitas a ficheiros de implementação mencionadas ao longo deste capítulo, agrupando-as por função arquitetural e componente do artefacto.",
            )
        ],
    )

    implementation_map_heading = find_heading(doc, "Mapa de Implementação")
    append_after(
        implementation_map_heading,
        [
            (
                "Normal",
                "Em complemento à Tabela 20, o Anexo B apresenta um mapeamento textual consolidado dos ficheiros de código e dos artefactos de avaliação citados nos Capítulos 4 e 5, facilitando a consulta cruzada entre descrição conceptual e implementação concreta.",
            )
        ],
    )

    evaluation_heading = find_heading(doc, "Critérios e Método de Avaliação")
    append_after(
        evaluation_heading,
        [
            (
                "Normal",
                "Os comandos, datasets e ficheiros de implementação diretamente associados à avaliação do router e do assistente encontram-se identificados no Anexo B, reforçando a reprodutibilidade e a auditabilidade da estratégia de avaliação adotada.",
            )
        ],
    )

    doc.add_page_break()
    doc.add_paragraph("Anexo B – Mapeamento das Referências ao Código e Artefactos de Avaliação", style="Heading 1")
    doc.add_paragraph(
        "O presente anexo reúne, de forma consolidada, os principais ficheiros de implementação e os artefactos de suporte à avaliação explicitamente referidos no corpo da dissertação. O objetivo é facilitar a inspeção técnica do artefacto e reforçar a correspondência entre a descrição conceptual apresentada no texto e a sua materialização concreta no código.",
        style="Normal",
    )
    for style, text in ANNEX_ITEMS:
        doc.add_paragraph(text, style=style)

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    main()
