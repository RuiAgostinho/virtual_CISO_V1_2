from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_BREAK
from docx.oxml import OxmlElement
from docx.shared import Pt
from docx.table import Table
from docx.text.paragraph import Paragraph


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "Relatorio_Rui_Agostinho_18_05_2026.docx"
OUTPUT = ROOT / "Relatorio_Rui_Agostinho_18_05_2026_revisto_cap4_5.docx"


def normalise(value: str) -> str:
    return " ".join(value.strip().split()).lower()


def find_paragraph(doc: Document, text: str) -> Paragraph:
    wanted = normalise(text)
    for paragraph in doc.paragraphs:
        if normalise(paragraph.text) == wanted:
            return paragraph
    raise ValueError(f"Paragraph not found: {text}")


def remove_body_range(doc: Document, start_text: str, end_text: str) -> Paragraph:
    start = find_paragraph(doc, start_text)
    end = find_paragraph(doc, end_text)
    body = doc._body._element
    children = list(body)
    start_idx = children.index(start._p)
    end_idx = children.index(end._p)
    if start_idx >= end_idx:
        raise ValueError("Invalid chapter range")
    for child in children[start_idx:end_idx]:
        body.remove(child)
    return end


def move_before(element, anchor: Paragraph) -> None:
    element.getparent().remove(element)
    anchor._p.addprevious(element)


def add_paragraph_before(doc: Document, anchor: Paragraph, text: str = "", style: str = "Normal") -> Paragraph:
    paragraph = doc.add_paragraph()
    if style:
        paragraph.style = style
    if text:
        paragraph.add_run(text)
    move_before(paragraph._p, anchor)
    return paragraph


def add_page_break_before(doc: Document, anchor: Paragraph) -> None:
    paragraph = add_paragraph_before(doc, anchor)
    paragraph.add_run().add_break(WD_BREAK.PAGE)


def add_bullets(doc: Document, anchor: Paragraph, items: list[str]) -> None:
    for item in items:
        add_paragraph_before(doc, anchor, item, "List Paragraph")


def add_table_before(
    doc: Document,
    anchor: Paragraph,
    caption: str,
    headers: list[str],
    rows: list[list[str]],
) -> Table:
    add_paragraph_before(doc, anchor, caption, "Caption")
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for idx, header in enumerate(headers):
        run = header_cells[idx].paragraphs[0].add_run(header)
        run.bold = True
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            cells[idx].text = value
    move_before(table._tbl, anchor)
    add_paragraph_before(doc, anchor)
    return table


def set_normal_font(doc: Document) -> None:
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:
            if run.font.size is None:
                run.font.size = Pt(11)


CHAPTER_4: list[tuple[str, object]] = [
    ("h1", "Conceção, Arquitetura e Desenvolvimento do Virtual CISO"),
    (
        "p",
        "A revisão sistemática apresentada no capítulo anterior evidenciou lacunas persistentes na articulação entre governação, gestão de risco e conformidade em cibersegurança. A literatura sobre GRC assinala, de forma recorrente, a fragmentação entre dados técnicos, requisitos normativos, mecanismos operacionais e evidências de auditoria. Esta fragmentação limita a capacidade do CISO para justificar decisões, demonstrar conformidade e priorizar investimento de forma rastreável. Em resposta a esse problema, o presente capítulo apresenta a conceção e o desenvolvimento do Virtual CISO, uma plataforma de apoio à decisão construída segundo a lógica de Design Science Research, conforme proposta por Hevner et al. (2004) e Peffers et al. (2007).",
    ),
    (
        "p",
        "O artefacto desenvolvido procura materializar uma ideia central: a conformidade não deve ser modelada a partir das frameworks externas, mas a partir dos controlos internos que a organização efetivamente decide implementar. As frameworks externas, como ISO/IEC 27001, ISO/IEC 27002, NIST CSF, NIS2, DORA, QNRC ou CIS Controls, passam assim a funcionar como camadas de mapeamento. Esta opção evita que a organização tenha de duplicar mecanismos e evidências para cada referencial, permitindo que um mesmo mecanismo interno, por exemplo MFA, segregação de funções, SIEM, EDR ou revisão periódica de acessos, contribua simultaneamente para vários referenciais.",
    ),
    (
        "p",
        "O Virtual CISO foi implementado como uma aplicação web modular, com backend em Django REST Framework, frontend em React com TypeScript, base de dados PostgreSQL e suporte a pesquisa vetorial com pgvector. A componente de IA utiliza modelos locais através de Ollama, integrados numa arquitetura de Retrieval-Augmented Generation (RAG), conceito proposto por Lewis et al. (2020), mas adaptado ao domínio GRC através de recuperação estruturada, recuperação semântica e normalização de fontes. A execução local dos modelos foi privilegiada por razões de confidencialidade e controlo, embora tenha implicações diretas no desempenho, avaliadas no Capítulo 5.",
    ),
    ("h2", "Requisitos e princípios arquiteturais"),
    (
        "p",
        "Os requisitos do Virtual CISO foram definidos a partir de três fontes: as lacunas identificadas na revisão sistemática, as necessidades práticas da função CISO e as exigências de rastreabilidade associadas a contextos regulados. A distinção entre requisitos funcionais e não funcionais segue a abordagem clássica de engenharia de software apresentada por Sommerville (2016), mas foi adaptada ao domínio específico de governação, risco e conformidade.",
    ),
    (
        "p",
        "Do ponto de vista funcional, a plataforma deve permitir gerir ativos, vulnerabilidades, riscos, controlos internos, frameworks externas, políticas, documentos de governação, mecanismos de implementação, tipos de evidência, evidências reais e recomendações assistidas por IA. Do ponto de vista não funcional, destacam-se explicabilidade, rastreabilidade, auditabilidade, controlo humano, segurança da informação, modularidade e capacidade de evolução incremental.",
    ),
    (
        "p",
        "O princípio arquitetural mais relevante é a separação entre catálogo interno e frameworks externas. No modelo anterior, o controlo estava diretamente associado a uma framework. Na arquitetura revista, o controlo interno passa a ser o centro do modelo e representa a forma como a organização descreve o seu próprio objetivo de controlo, de forma agnóstica face aos referenciais. A partir desse catálogo interno, são estabelecidos mapeamentos validados para controlos externos. Esta decisão aumenta a coerência da plataforma, reduz duplicação de evidências e permite calcular conformidade por propagação.",
    ),
    (
        "table",
        {
            "caption": "Quadro 4.1 - Princípios de conceção adotados no Virtual CISO",
            "headers": ["Princípio", "Aplicação no artefacto", "Justificação"],
            "rows": [
                [
                    "Internal-first",
                    "InternalControl é o centro do modelo; frameworks externas são mapeamentos.",
                    "Evita duplicação de controlos e evidencia a postura real da organização.",
                ],
                [
                    "Human-in-the-loop",
                    "Mappings, políticas, mecanismos, evidências e recomendações exigem validação humana.",
                    "Mantém o CISO responsável pela decisão e reduz automatismos opacos.",
                ],
                [
                    "Rastreabilidade por desenho",
                    "TraceabilityService e SourceNormalizer expõem relações, fontes e caminhos de decisão.",
                    "Suporta auditoria, revisão e explicação das recomendações.",
                ],
                [
                    "Compatibilidade incremental",
                    "Modelos legacy são preservados e ligados à nova camada por migrações e adapters.",
                    "Permite evoluir a aplicação sem perda de dados nem quebra de endpoints existentes.",
                ],
                [
                    "IA ancorada em contexto",
                    "LLMRouter, GovernanceContextAdapter, RAG e fontes normalizadas condicionam a resposta.",
                    "Reduz respostas genéricas e privilegia informação interna antes da externa.",
                ],
            ],
        },
    ),
    ("h2", "Arquitetura geral do sistema"),
    (
        "p",
        "A arquitetura do Virtual CISO segue uma organização em camadas, adequada a sistemas complexos que exigem separação de responsabilidades, evolução incremental e capacidade de integração, conforme defendido por Bass et al. (2021). A opção por uma arquitetura modular dentro de uma aplicação coerente, em vez de uma decomposição prematura em microsserviços, justifica-se pela natureza exploratória do artefacto e pela necessidade de preservar relações ricas entre entidades de governação, risco e conformidade.",
    ),
    (
        "table",
        {
            "caption": "Quadro 4.2 - Camadas arquiteturais do Virtual CISO",
            "headers": ["Camada", "Componentes principais", "Responsabilidade"],
            "rows": [
                [
                    "Apresentação",
                    "React, TypeScript, Vite, páginas de governance, risk, assets, compliance, admin e IA.",
                    "Disponibilizar uma experiência de trabalho para o CISO, com dashboards, wizards e páginas de detalhe.",
                ],
                [
                    "API",
                    "Django REST Framework, autenticação JWT, viewsets e serializers.",
                    "Expor endpoints para entidades GRC, mappings, scoring, traceability e assistente.",
                ],
                [
                    "Domínio",
                    "Modelos Django: InternalControl, Policy, GovernanceDocument, Mechanism, EvidenceItem, Framework e Risk.",
                    "Representar entidades internas, relações, estados, validação humana e histórico.",
                ],
                [
                    "Serviços analíticos",
                    "RiskEngine, VulnerabilityPrioritizationService, CompliancePropagationEngine e TraceabilityService.",
                    "Calcular risco, prioridade, conformidade por propagação, gaps e rastreabilidade.",
                ],
                [
                    "IA e conhecimento",
                    "LLMRouter, QueryOrchestrator, StructuredQueryService, SemanticRetrievalService, PromptBuilder, GovernanceContextAdapter e pgvector.",
                    "Responder a perguntas, recuperar contexto, gerar recomendações e apresentar fontes.",
                ],
                [
                    "Integrações",
                    "Wazuh, Nmap, EPSS, NIST/NVD, CISA KEV e mecanismos de ingestão.",
                    "Enriquecer a base operacional com ativos, vulnerabilidades, exposição e inteligência de ameaças.",
                ],
            ],
        },
    ),
    (
        "p",
        "A camada de apresentação foi desenhada para suportar tanto leitura executiva como trabalho operacional. O Mission Control apresenta indicadores consolidados; as páginas de políticas, documentos, mecanismos e evidências permitem navegação detalhada; os wizards reduzem a complexidade dos fluxos de criação; e o Mapping Review Dashboard concentra a validação humana dos mapeamentos críticos. Esta composição procura responder à crítica identificada na literatura a dashboards meramente descritivos, aproximando a interface de um verdadeiro sistema de apoio à decisão.",
    ),
    ("h2", "Modelo de dados transversal e catálogo interno de controlos"),
    (
        "p",
        "A principal evolução arquitetural introduzida no desenvolvimento do Virtual CISO foi a criação de uma camada transversal de governação centrada no catálogo interno de controlos. O modelo Control existente foi preservado como representação de controlos externos pertencentes a frameworks. Sobre esse modelo foi criada a entidade InternalControl, que representa o controlo definido pela organização, independente de uma norma específica.",
    ),
    (
        "p",
        "A entidade InternalControl possui código, título, descrição, domínio, objetivo, declaração de risco, responsável, criticidade, estado, origem e ligação opcional ao controlo legacy. Esta ligação permite rastrear a origem de controlos migrados, sem transformar a framework externa no centro da arquitetura. A associação entre controlos internos e controlos externos é realizada por InternalControlFrameworkMapping, que inclui tipo de relação, percentagem de cobertura, rationale, origem do mapping, estado de validação e score de confiança.",
    ),
    (
        "p",
        "Este desenho permite responder a um problema prático comum em GRC: o mesmo controlo organizacional aparece com formulações diferentes em múltiplos referenciais. Por exemplo, um controlo interno relativo a autenticação forte pode mapear para controlos ISO/IEC 27002, subcategorias NIST CSF, requisitos NIS2 e práticas CIS Controls. Ao centralizar a implementação no controlo interno e mapear para os referenciais externos, a plataforma consegue calcular impacto e conformidade por propagação, sem exigir que o CISO repita a mesma evidência para cada framework.",
    ),
    ("h3", "Políticas, documentos e estrutura normativa"),
    (
        "p",
        "A camada documental foi alargada para distinguir a política enquanto objeto de governação e os documentos que dela podem derivar. A entidade Policy continua disponível para compatibilidade, mas passa a articular-se com InternalControl através de PolicyInternalControl. Em paralelo, GovernanceDocument representa documentos de governação mais genéricos, incluindo políticas, normas, procedimentos, guidelines, regulamentos técnicos e runbooks. GovernanceDocumentSection permite construir capítulos e subcapítulos, aproximando a edição de políticas de um processo real de redação normativa.",
    ),
    (
        "p",
        "Esta estrutura é particularmente relevante porque a governação de segurança não se limita a declarar controlos. Um CISO necessita de redigir políticas, decompor responsabilidades, criar normas subordinadas, operacionalizar procedimentos e, em alguns casos, disponibilizar runbooks com passos executáveis. A associação GovernanceDocumentControl liga documentos a controlos internos e permite indicar se o documento define, implementa, operacionaliza ou evidencia determinado controlo.",
    ),
    ("h3", "Mecanismos, estado operacional e evidência"),
    (
        "p",
        "Entre o controlo interno e a evidência foi introduzida uma camada explícita de mecanismos de implementação. A entidade Mechanism funciona como catálogo reutilizável de mecanismos, como MFA, RBAC, PAM, SIEM, EDR/XDR, WAF, backups, encriptação, segmentação de rede, revisão periódica de acessos ou offboarding. A associação InternalControlMechanism liga esses mecanismos aos controlos internos, indicando tipo de relação, peso de contribuição, obrigatoriedade, rationale, estado de validação e estado operacional de implementação.",
    ),
    (
        "p",
        "O estado operacional implementation_status foi colocado em InternalControlMechanism e não em Mechanism, porque o mesmo mecanismo pode estar implementado para um controlo e apenas planeado para outro. Os estados suportados são not_implemented, planned, partially_implemented, implemented, implemented_evidenced e not_applicable. Esta decisão melhora o realismo do scoring, uma vez que o cumprimento de um controlo deixa de depender apenas da existência abstrata de um mecanismo e passa a considerar a sua implementação naquele contexto específico.",
    ),
    (
        "p",
        "A evidência foi igualmente redesenhada em duas dimensões. A primeira corresponde aos tipos de evidência esperada por mecanismo, representados por MechanismEvidenceRequirement. Estes tipos ajudam o CISO a perceber que provas deve recolher para demonstrar que um mecanismo está implementado. A segunda corresponde às evidências reais, representadas por EvidenceItem e associadas a entidades através de EvidenceLink. Assim, a plataforma distingue claramente entre aquilo que deve ser recolhido e aquilo que já foi efetivamente anexado, validado ou expirou.",
    ),
    ("h2", "Scoring de conformidade por propagação"),
    (
        "p",
        "A nova arquitetura permitiu criar um motor de scoring separado da lógica legacy: o CompliancePropagationEngine. Este serviço calcula conformidade a partir dos mecanismos e evidências, propagando depois o resultado para controlos internos, políticas, documentos de governação, controlos externos e frameworks. O cálculo oficial considera apenas mapeamentos aprovados; o modo de simulação permite incluir elementos pending_review; e elementos rejected ou deprecated nunca são considerados ativos.",
    ),
    (
        "p",
        "O score de um mecanismo no contexto de um controlo interno resulta do implementation_status e da existência de evidência válida. Um mecanismo not_implemented tem score 0; planned corresponde a 20; partially_implemented a 40; implemented a 70; e implemented_evidenced só atinge 100 se existir EvidenceItem válida associada por EvidenceLink aprovado. Evidências expiradas, rejeitadas ou obsoletas não contam para o score oficial. Esta regra torna explícita a diferença entre afirmar que um mecanismo existe e demonstrar que ele está efetivamente implementado.",
    ),
    (
        "p",
        "O score de InternalControl é calculado como média ponderada dos mecanismos associados, usando contribution_weight, mandatory e validation_status. A existência de um mecanismo obrigatório não implementado impede que o controlo seja considerado plenamente conforme, mesmo que a média agregada seja elevada. O score de Policy e GovernanceDocument resulta dos controlos internos associados, enquanto o score de Control externo é obtido pelos InternalControls mapeados através de InternalControlFrameworkMapping. Por fim, o score de Framework é agregado a partir dos controlos externos avaliados, acompanhado por uma métrica de coverage para evitar esconder falta de cobertura.",
    ),
    (
        "table",
        {
            "caption": "Quadro 4.3 - Propagação de conformidade na nova arquitetura",
            "headers": ["Nível", "Base de cálculo", "Resultado esperado"],
            "rows": [
                [
                    "Mechanism em contexto",
                    "InternalControlMechanism.implementation_status e EvidenceLinks válidos.",
                    "Score operacional do mecanismo para um controlo concreto.",
                ],
                [
                    "InternalControl",
                    "Média ponderada dos mecanismos aprovados associados ao controlo.",
                    "Score interno, gaps e estado de conformidade.",
                ],
                [
                    "Policy / GovernanceDocument",
                    "Controlos internos associados por PolicyInternalControl ou GovernanceDocumentControl.",
                    "Score documental e lacunas de governação.",
                ],
                [
                    "Control externo",
                    "InternalControls mapeados e coverage_percentage.",
                    "Score por controlo de framework.",
                ],
                [
                    "Framework",
                    "Scores dos controlos externos avaliados e cobertura global.",
                    "Score avaliado, coverage e principais gaps.",
                ],
            ],
        },
    ),
    ("h2", "Rastreabilidade transversal"),
    (
        "p",
        "A rastreabilidade foi implementada como uma camada própria através do TraceabilityService. Este serviço permite consultar, a partir de uma Policy, GovernanceDocument, InternalControl, Mechanism, EvidenceItem, Framework ou Control externo, as relações diretas e indiretas relevantes para governação e conformidade. A resposta inclui root, relationships, active_mappings, inactive_mappings, scores, evidence, gaps e metadata, de acordo com os parâmetros de modo, profundidade e inclusão de detalhes.",
    ),
    (
        "p",
        "Esta capacidade é central para a natureza auditável do artefacto. A partir de uma política, o CISO pode identificar os controlos internos associados, os mecanismos que operacionalizam esses controlos, os tipos de evidência esperada, as evidências reais recolhidas, os controlos externos impactados e as frameworks afetadas. No sentido inverso, a partir de uma framework, é possível perceber que controlos internos, mecanismos, políticas e evidências sustentam o score apresentado.",
    ),
    ("h2", "IA generativa, RAG híbrido e adapter internal-first"),
    (
        "p",
        "A componente de IA foi concebida para apoiar aconselhamento, síntese e redação, mas não para substituir a decisão do CISO. O pipeline mantém o LLMRouter como componente de classificação do pedido, sem o alterar para acomodar a nova arquitetura. A integração dos novos dados de governance é realizada por adapter, através do GovernanceContextAdapter, ligado ao QueryOrchestrator e ao StructuredQueryService. Esta opção preserva a responsabilidade do router e adiciona contexto interno antes da geração da resposta.",
    ),
    (
        "p",
        "O funcionamento internal-first significa que o assistente procura primeiro controlos internos, políticas internas, documentos de governação, mecanismos, evidências, scoring e traceability. Só depois recorre a controlos externos de frameworks como contexto complementar. Esta decisão é coerente com a arquitetura de governação adotada: o conhecimento organizacional interno é a fonte primária de decisão; as frameworks externas servem para enquadrar, comparar e demonstrar conformidade.",
    ),
    (
        "p",
        "O processo de RAG articula QueryOrchestrator, StructuredQueryService, SemanticRetrievalService, PromptBuilder, SourceNormalizer e OllamaClient. Sempre que possível, perguntas factuais ou agregadas são respondidas por consulta estruturada. Quando a resposta exige contexto documental ou semântico, o sistema recupera KnowledgeChunks indexados em pgvector. A ingestão de conteúdos internos gera chunks e embeddings para controlos internos, políticas, documentos, secções, mecanismos, evidências, tipos de evidência e mappings. A administração da reindexação permite atualizar o índice quando são criadas ou alteradas entidades relevantes.",
    ),
    (
        "p",
        "A integração do assistente na edição de políticas materializa este princípio. Ao redigir uma política, o CISO pode solicitar revisão, recomendações, identificação de gaps ou proposta de redação. O assistente constrói a resposta com base no conteúdo da política, controlos internos associados, mecanismos, evidências esperadas, frameworks impactadas e histórico persistente de recomendações. As respostas são guardadas, podem ser consultadas posteriormente e incluem fontes visíveis, reforçando a rastreabilidade.",
    ),
    ("h2", "Interface, wizards e experiência do CISO"),
    (
        "p",
        "A evolução da interface acompanhou a evolução do modelo de dados. Foram criadas páginas e wizards para reduzir a necessidade de manipulação direta de entidades técnicas. O PolicyWizard permite criar uma política e associá-la a controlos internos; o GovernanceDocumentWizard cria normas, procedimentos, guidelines, regulamentos técnicos e runbooks; o MechanismWizard permite criar ou reutilizar mecanismos e associá-los a controlos internos; o EvidenceWizard permite criar evidências reais e associá-las a entidades; e o FrameworkMappingWizard suporta o mapeamento entre controlos internos e controlos externos.",
    ),
    (
        "p",
        "O Mapping Review Dashboard concentra a validação humana dos mapeamentos. Nesta área, o CISO consegue visualizar mapeamentos pending_review, draft, approved, rejected ou deprecated, filtrar por entidade, aprovar, rejeitar, marcar como obsoleto e consultar impacto por traceability. Esta etapa explícita de validação reforça a defensabilidade da solução, porque separa sugestões, migrações e relações técnicas de uma decisão formal validada por responsável humano.",
    ),
    (
        "p",
        "As páginas de detalhe de políticas foram também ajustadas para permitir leitura integral do conteúdo documental, edição estruturada por capítulos e subcapítulos, onboarding da política, seleção de controlos internos, mecanismos associados, tipos de evidência esperada, evidências reais recolhidas, exportação para PDF e análise de gaps por IA. Esta experiência aproxima a aplicação do trabalho real de um CISO: redigir, relacionar, evidenciar, validar e comunicar.",
    ),
    (
        "table",
        {
            "caption": "Quadro 4.4 - Mapa de implementação das principais capacidades revistas",
            "headers": ["Capacidade", "Artefactos principais", "Localização indicativa"],
            "rows": [
                [
                    "Catálogo interno de controlos",
                    "InternalControl, InternalControlFrameworkMapping",
                    "backend/governance/models/",
                ],
                [
                    "Políticas e documentos",
                    "PolicyInternalControl, GovernanceDocument, GovernanceDocumentSection, RunbookStep",
                    "backend/governance/models/ e frontend/src/pages/governance/",
                ],
                [
                    "Mecanismos e evidências",
                    "InternalControlMechanism, MechanismEvidenceRequirement, EvidenceItem, EvidenceLink",
                    "backend/governance/models/ e frontend/src/lib/mappingReviewApi.ts",
                ],
                [
                    "Scoring por propagação",
                    "CompliancePropagationEngine, CompliancePropagationResult",
                    "backend/governance/services/compliance_propagation_engine.py",
                ],
                [
                    "Rastreabilidade",
                    "TraceabilityService e endpoints /api/governance/traceability/",
                    "backend/governance/services/traceability_service.py",
                ],
                [
                    "RAG internal-first",
                    "GovernanceContextAdapter, QueryOrchestrator, StructuredQueryService, SemanticRetrievalService",
                    "backend/ciso_assistant/services/",
                ],
                [
                    "Validação de mappings",
                    "Mapping Review Dashboard e actions approve/reject/deprecated",
                    "frontend/src/pages/governance/MappingReview.tsx",
                ],
            ],
        },
    ),
    ("h2", "Síntese do capítulo"),
    (
        "p",
        "A arquitetura apresentada neste capítulo evolui o Virtual CISO de uma plataforma de GRC com suporte multi-framework para um artefacto transversal, centrado no conhecimento interno da organização. A distinção entre controlos internos e controlos externos, a associação de políticas e documentos a controlos internos, a ligação entre mecanismos e evidências, o scoring por propagação, a rastreabilidade ponta-a-ponta e a integração RAG internal-first constituem os principais contributos técnicos desta fase de desenvolvimento. O capítulo seguinte demonstra e avalia estas capacidades, incluindo o impacto da execução local dos LLMs no desempenho observado.",
    ),
]


CHAPTER_5: list[tuple[str, object]] = [
    ("h1", "Demonstração e Avaliação do Virtual CISO"),
    (
        "p",
        "O presente capítulo apresenta a demonstração e avaliação do Virtual CISO, em coerência com as fases de demonstração e avaliação da metodologia Design Science Research. A avaliação procura verificar se o artefacto desenvolvido responde ao problema identificado: apoiar o CISO na articulação entre governação, risco e conformidade, através de informação estruturada, mecanismos analíticos, rastreabilidade e aconselhamento assistido por IA.",
    ),
    (
        "p",
        "A avaliação não incide apenas sobre o assistente de IA. Embora a componente conversacional seja relevante, o valor do Virtual CISO resulta da integração entre catálogo interno de controlos, políticas, documentos, mecanismos, evidências, frameworks, risco, scoring e traceability. Assim, a demonstração foi organizada em torno de fluxos completos de trabalho, mais próximos da atividade real de um CISO do que de testes isolados de componentes.",
    ),
    ("h2", "Estratégia de demonstração e avaliação"),
    (
        "p",
        "A demonstração foi conduzida sobre um cenário representativo de organização com necessidades de governação e conformidade em cibersegurança. O cenário permite demonstrar a criação de políticas, a associação de controlos internos, a definição de mecanismos, a recolha de evidências, o mapeamento para frameworks externas, o cálculo de conformidade por propagação e a utilização do assistente para análise de gaps e recomendações.",
    ),
    (
        "p",
        "A avaliação combina observação funcional, análise de cobertura de requisitos, testes técnicos e medições de desempenho. Em particular, foram considerados: cobertura funcional dos fluxos principais; coerência do modelo internal-first; qualidade da rastreabilidade; comportamento do CompliancePropagationEngine; desempenho do LLMRouter; comportamento do RAG internal-first; latência das respostas de IA; e limitações decorrentes da execução em ambiente local com recursos computacionais reduzidos.",
    ),
    (
        "table",
        {
            "caption": "Quadro 5.1 - Dimensões de avaliação consideradas",
            "headers": ["Dimensão", "Questão de avaliação", "Instrumento"],
            "rows": [
                [
                    "Funcional",
                    "O artefacto permite executar os fluxos centrais de governação e conformidade?",
                    "Walkthrough funcional, wizards, páginas de detalhe e Mapping Review.",
                ],
                [
                    "Arquitetural",
                    "A camada internal-first reduz duplicação e preserva compatibilidade legacy?",
                    "Análise do modelo de dados, endpoints e migrações incrementais.",
                ],
                [
                    "Analítica",
                    "O scoring por propagação reflete mecanismos, evidências e validação humana?",
                    "Testes do CompliancePropagationEngine e cenários de score.",
                ],
                [
                    "Rastreabilidade",
                    "É possível navegar da política até mecanismos, evidências e frameworks, e no sentido inverso?",
                    "Traceability API e painéis de detalhe.",
                ],
                [
                    "IA/RAG",
                    "O assistente usa primeiro contexto interno e apresenta fontes úteis?",
                    "Dataset de perguntas, histórico de recomendações e avaliação RAG.",
                ],
                [
                    "Desempenho",
                    "A latência é aceitável no ambiente local utilizado?",
                    "Medição de tempos por tipo de pergunta e por modelo local.",
                ],
            ],
        },
    ),
    ("h2", "Cenário de demonstração"),
    (
        "p",
        "O cenário de demonstração corresponde a uma organização que pretende formalizar e operacionalizar a sua política de segurança da informação. O CISO necessita de redigir a política, associá-la a controlos internos, criar documentos subordinados, identificar mecanismos de implementação, definir evidências esperadas, recolher evidências reais, mapear controlos internos para frameworks externas e avaliar a conformidade resultante.",
    ),
    (
        "p",
        "O fluxo demonstrado segue a cadeia: criar política, associar controlos internos, criar documento de governação, associar mecanismos, criar ou associar evidências, mapear frameworks, validar score e consultar traceability. Este fluxo foi escolhido por evidenciar a proposta central do artefacto: a conformidade passa a ser demonstrada pela relação entre controlo interno, mecanismo, evidência e mapping para frameworks, e não por checklists isoladas de cada referencial.",
    ),
    (
        "table",
        {
            "caption": "Quadro 5.2 - Fluxo de demonstração da camada transversal de governação",
            "headers": ["Etapa", "Ação demonstrada", "Resultado observado / a registar"],
            "rows": [
                [
                    "1. Política",
                    "Criar ou editar uma política de segurança com capítulos e subcapítulos.",
                    "[Inserir política usada, número de capítulos e estado final.]",
                ],
                [
                    "2. Controlos internos",
                    "Associar controlos internos agnósticos de frameworks.",
                    "[Inserir número de controlos internos associados.]",
                ],
                [
                    "3. Documento",
                    "Criar norma, procedimento, guideline ou runbook subordinado.",
                    "[Inserir tipo de documento criado e relação hierárquica.]",
                ],
                [
                    "4. Mecanismos",
                    "Associar mecanismos reutilizáveis aos controlos internos.",
                    "[Inserir número de mecanismos associados e estados de implementação.]",
                ],
                [
                    "5. Evidências",
                    "Definir tipos de evidência esperada e anexar evidências reais.",
                    "[Inserir número de tipos esperados, evidências reais e estado de validação.]",
                ],
                [
                    "6. Framework mapping",
                    "Mapear controlos internos para controlos externos.",
                    "[Inserir frameworks impactadas e número de mappings aprovados.]",
                ],
                [
                    "7. Score",
                    "Calcular conformidade por propagação.",
                    "[Inserir score official, score simulation e coverage.]",
                ],
                [
                    "8. Traceability",
                    "Consultar cadeia Policy -> InternalControl -> Mechanism -> Evidence -> Framework.",
                    "[Inserir observações sobre relações devolvidas e lacunas identificadas.]",
                ],
            ],
        },
    ),
    ("h2", "Demonstração funcional dos principais fluxos"),
    ("h3", "Criação e redação de políticas"),
    (
        "p",
        "A criação e edição de políticas foi demonstrada através da página dedicada de detalhe e edição documental. O editor permite estruturar a política em capítulos e subcapítulos, alterar a ordem das secções, redigir conteúdo livre e manter metadados como owner, accountable, unidade responsável, estado, versão e datas de revisão. Esta capacidade é relevante porque o CISO não trabalha apenas com entidades técnicas: uma parte essencial da função consiste em produzir documentos normativos claros, aprováveis e auditáveis.",
    ),
    ("h3", "Associação de controlos internos"),
    (
        "p",
        "A associação de controlos internos foi realizada no onboarding da política e no PolicyWizard. A interface apresenta primeiro controlos internos, e só recorre a controlos externos como contexto complementar. Esta alteração é coerente com a arquitetura internal-first e evita confusão entre o catálogo organizacional e controlos de frameworks externas. Cada associação PolicyInternalControl mantém applicability, rationale, origem e estado de validação.",
    ),
    ("h3", "Mecanismos e tipos de evidência esperada"),
    (
        "p",
        "A demonstração confirmou a importância de distinguir mecanismos de implementação e tipos de evidência esperada. Quando um controlo interno é selecionado, a plataforma apresenta os mecanismos associados a esse controlo. Para cada mecanismo, são apresentados tipos de evidência recomendados, como configuração aprovada, log de execução, relatório de revisão, ticket de aprovação, ata, exportação técnica ou evidência de monitorização. Estes tipos não são, por si só, evidência real; funcionam como checklist operacional para o CISO saber o que tem de recolher.",
    ),
    (
        "p",
        "A evidência real é registada como EvidenceItem e associada através de EvidenceLink. Esta separação torna o modelo mais defensável: o sistema pode indicar que determinado mecanismo exige evidência, mas só considera o mecanismo plenamente evidenciado quando uma evidência real válida estiver associada, aprovada e dentro do prazo de validade.",
    ),
    ("h3", "Mapping Review e validação humana"),
    (
        "p",
        "O Mapping Review Dashboard foi avaliado como ponto explícito de validação. Mappings criados automaticamente, migrados ou sugeridos podem permanecer em draft ou pending_review até serem aprovados. A aprovação, rejeição ou marcação como deprecated ficam separadas do simples ato de criação. Esta etapa é fundamental para a defensabilidade da solução, porque permite demonstrar que a relação entre política, controlo, mecanismo, evidência e framework foi validada por uma pessoa responsável.",
    ),
    ("h3", "Scoring e traceability"),
    (
        "p",
        "A conformidade por propagação foi demonstrada através dos endpoints de CompliancePropagation e da Traceability API. O score de uma política passou a depender dos controlos internos associados; estes dependem dos mecanismos; e os mecanismos dependem do seu estado operacional e da existência de evidência válida. A Traceability API permite confirmar o caminho usado no cálculo, bem como identificar lacunas, mapeamentos pendentes e entidades sem evidência.",
    ),
    ("h2", "Avaliação da componente de IA e RAG"),
    (
        "p",
        "A componente de IA foi avaliada em duas dimensões: routing e qualidade contextual das respostas. O LLMRouter classifica a pergunta e decide se deve usar consulta estruturada, priorização determinística, recuperação RAG ou geração direta. A integração da nova camada de governance foi realizada sem alterar o router, através do GovernanceContextAdapter. Este adapter recupera primeiro contexto interno, como InternalControls, políticas, documentos, mecanismos, evidências, mappings e resultados de scoring, e só depois utiliza controlos externos como contexto de apoio.",
    ),
    (
        "p",
        "Este comportamento é essencial para o objetivo do trabalho. Quando o CISO pergunta sobre uma política, o assistente deve aconselhar com base na política interna, nos controlos internos associados, nos mecanismos exigidos, nas evidências esperadas e nas frameworks impactadas. Uma resposta que começasse por listar controlos externos genéricos, sem considerar a realidade interna, seria menos útil e menos defensável.",
    ),
    (
        "p",
        "Foram também demonstradas recomendações persistentes na página de histórico de recomendações do assistente. As análises de gaps e recomendações associadas a uma política ficam armazenadas, podem ser filtradas por contexto e podem ser usadas como apoio à decisão formal. Esta persistência permite passar de uma interação conversacional efémera para um registo auditável de aconselhamento.",
    ),
    (
        "table",
        {
            "caption": "Quadro 5.3 - Testes recomendados para avaliação da IA e RAG",
            "headers": ["Teste", "Procedimento", "Resultado a preencher"],
            "rows": [
                [
                    "Routing do LLMRouter",
                    "Executar dataset docs/llmrouter_test_queries.json e comparar task_type e needs_rag.",
                    "[N.º de perguntas]; [acurácia task_type]; [acurácia needs_rag]; [erros observados].",
                ],
                [
                    "RAG internal-first",
                    "Executar perguntas sobre políticas e controlos internos e verificar se as primeiras fontes são internas.",
                    "[% respostas com fonte interna no top-3]; [exemplos de fontes corretas/incorretas].",
                ],
                [
                    "Aconselhamento de política",
                    "Pedir revisão completa, redação normativa e gaps para uma política concreta.",
                    "[tempo médio]; [qualidade percebida]; [n.º fontes]; [principais gaps identificados].",
                ],
                [
                    "Histórico persistente",
                    "Guardar recomendações e verificar apresentação em Recommendation History e Mission Control.",
                    "[N.º recomendações guardadas]; [filtros testados]; [observações].",
                ],
                [
                    "Reindexação",
                    "Criar/alterar controlos, políticas, mecanismos e evidências; executar reindexação pelo frontend/admin.",
                    "[n.º chunks criados]; [n.º embeddings gerados]; [tempo de reindexação]; [falhas].",
                ],
                [
                    "Falha do LLM local",
                    "Parar o serviço LLM ou simular timeout e verificar fallback/mensagem ao utilizador.",
                    "[comportamento observado]; [mensagem apresentada]; [impacto na experiência].",
                ],
            ],
        },
    ),
    ("h2", "Avaliação do desempenho e limitação do LLM local"),
    (
        "p",
        "Uma dimensão crítica da avaliação diz respeito ao tempo de resposta da IA. O Virtual CISO executa modelos de linguagem localmente, através de Ollama, em vez de recorrer a APIs cloud. Esta decisão tem vantagens relevantes em confidencialidade, controlo dos dados e independência operacional, mas introduz uma limitação clara: a latência das respostas depende fortemente dos recursos computacionais disponíveis, nomeadamente CPU, GPU, RAM, velocidade de disco, modelo selecionado e tamanho do contexto enviado ao LLM.",
    ),
    (
        "p",
        "Nos testes exploratórios realizados durante o desenvolvimento, observou-se que perguntas estruturadas respondidas por ORM são rápidas, porque não exigem geração por LLM. Em contraste, perguntas com RAG e aconselhamento de política são mais lentas, pois envolvem classificação, recuperação de contexto, construção de prompt e geração local. Esta diferença deve ser explicitada na avaliação para evitar interpretar a lentidão como falha funcional: trata-se de uma consequência da opção arquitetural por execução local em hardware limitado.",
    ),
    (
        "p",
        "A avaliação deve, por isso, separar qualidade funcional e desempenho. Um modelo local pequeno pode produzir respostas úteis para aconselhamento, mas demorar dezenas de segundos ou mais quando o contexto é extenso. Modelos maiores tendem a melhorar a qualidade da redação, mas aumentam tempo de geração e consumo de memória. A solução é adequada como prova de conceito e como opção privacy-first, mas exigiria hardware superior, quantização adequada, streaming de resposta, cache, workers assíncronos ou eventual integração opcional com modelos externos para cenários produtivos.",
    ),
    (
        "table",
        {
            "caption": "Quadro 5.4 - Plano de medição de latência da componente de IA",
            "headers": ["Tipo de pergunta", "Modelo / contexto", "Métricas a registar"],
            "rows": [
                [
                    "Consulta estruturada sem LLM",
                    "ORM / StructuredQueryService",
                    "[tempo médio]; [p95]; [n.º execuções]; [observações].",
                ],
                [
                    "Pergunta geral sem RAG",
                    "Modelo local selecionado pelo LLMRouter",
                    "[modelo]; [tempo médio]; [tokens aproximados]; [qualidade percebida].",
                ],
                [
                    "Pergunta com RAG",
                    "Top-K chunks internos + geração local",
                    "[n.º chunks]; [tempo retrieval]; [tempo geração]; [tempo total].",
                ],
                [
                    "Revisão completa de política",
                    "Política + controlos internos + mecanismos + evidências + frameworks",
                    "[tempo total]; [n.º fontes]; [modelo]; [memória/CPU/GPU se disponível].",
                ],
                [
                    "Redação normativa",
                    "Pedido de exemplo de redação com contexto documental",
                    "[tempo total]; [extensão da resposta]; [necessidade de ajuste manual].",
                ],
                [
                    "Análise de gaps IA",
                    "Assistente integrado na página de detalhe da política",
                    "[tempo total]; [n.º gaps]; [n.º recomendações]; [resposta guardada no histórico?].",
                ],
            ],
        },
    ),
    ("h2", "Avaliação funcional da nova camada de governação"),
    (
        "p",
        "A avaliação funcional da camada internal-first deve verificar se a plataforma permite criar, consultar, relacionar e validar entidades sem depender diretamente das frameworks externas. A cadeia central Policy -> InternalControl -> Mechanism -> EvidenceItem -> FrameworkControl -> Framework foi demonstrada através das páginas de políticas, onboarding, Mapping Review, páginas de mecanismos, evidências e Traceability API.",
    ),
    (
        "table",
        {
            "caption": "Quadro 5.5 - Testes funcionais recomendados para a camada de governação",
            "headers": ["Teste", "Critério de sucesso", "Resultado a preencher"],
            "rows": [
                [
                    "Criar política com capítulos",
                    "A política guarda título, metadados, capítulos e subcapítulos.",
                    "[OK/NOK]; [observações].",
                ],
                [
                    "Associar controlos internos",
                    "A lista apresenta primeiro controlos internos e cria PolicyInternalControl em draft.",
                    "[OK/NOK]; [n.º controlos associados].",
                ],
                [
                    "Associar mecanismos",
                    "Mecanismos aparecem no contexto dos controlos selecionados e permitem definir implementation_status.",
                    "[OK/NOK]; [n.º mecanismos associados].",
                ],
                [
                    "Definir tipos de evidência",
                    "Cada mecanismo apresenta tipos de evidência esperada relevantes.",
                    "[OK/NOK]; [n.º tipos por mecanismo].",
                ],
                [
                    "Anexar evidência real",
                    "EvidenceItem fica ligado ao mecanismo por EvidenceLink e visível na página de evidências.",
                    "[OK/NOK]; [estado da evidência].",
                ],
                [
                    "Validar Mapping Review",
                    "Mappings podem ser aprovados, rejeitados e marcados como deprecated.",
                    "[OK/NOK]; [n.º mappings validados].",
                ],
                [
                    "Calcular score",
                    "CompliancePropagation devolve score, status, gaps e coverage.",
                    "[score official]; [score simulation]; [coverage].",
                ],
                [
                    "Consultar traceability",
                    "Traceability devolve políticas, controlos, mecanismos, evidências e frameworks impactadas.",
                    "[OK/NOK]; [lacunas detetadas].",
                ],
            ],
        },
    ),
    ("h2", "Discussão dos resultados"),
    (
        "p",
        "Os resultados observados sustentam a viabilidade da abordagem proposta. A nova arquitetura permite representar a realidade interna da organização antes de a projetar sobre frameworks externas. Esta opção aproxima o artefacto de uma prática de governação mais madura, na qual o CISO gere objetivos de controlo, mecanismos, evidências e responsabilidades, e só depois demonstra como estes se alinham com referenciais normativos.",
    ),
    (
        "p",
        "A camada de mecanismos e evidências mostrou-se particularmente relevante. Ao separar mecanismos de implementação, tipos de evidência esperada e evidências reais recolhidas, o sistema permite distinguir intenção, operacionalização e prova. Esta distinção é essencial em auditoria: não basta afirmar que existe um controlo; é necessário indicar que mecanismo o implementa e que evidência demonstra a sua execução.",
    ),
    (
        "p",
        "A componente de IA acrescenta valor quando atua sobre esta estrutura. As respostas deixam de ser recomendações genéricas sobre boas práticas e passam a considerar controlos internos, documentos, mecanismos, evidências, gaps e frameworks impactadas. Ainda assim, a avaliação mostra que a IA deve ser entendida como apoio, não como autoridade final. O CISO continua responsável por aprovar mappings, validar recomendações e decidir que ações executar.",
    ),
    (
        "p",
        "O principal compromisso identificado situa-se no desempenho. A execução local dos LLMs reforça privacidade e controlo, mas torna as respostas mais lentas em hardware limitado. Esta limitação não invalida a proposta, mas deve ser apresentada com transparência: a arquitetura é defensável para prova de conceito e ambientes que privilegiam confidencialidade; para produção, deverá ser acompanhada por otimizações de infraestrutura e processamento assíncrono.",
    ),
    ("h2", "Limitações da avaliação"),
    (
        "p",
        "A primeira limitação decorre da natureza de prova de conceito. O cenário utilizado é representativo, mas não corresponde ainda a uma avaliação longitudinal em organização real. Como tal, os resultados demonstram viabilidade conceptual e funcional, mas não impacto organizacional mensurado ao longo do tempo.",
    ),
    (
        "p",
        "A segunda limitação prende-se com a ausência de avaliação sistemática por utilizadores finais. A interação de CISOs, responsáveis de conformidade, auditores e equipas técnicas permitiria avaliar usabilidade, confiança, utilidade percebida e adequação dos fluxos ao trabalho quotidiano.",
    ),
    (
        "p",
        "A terceira limitação está relacionada com a componente de IA. Os datasets de teste são adequados para validação inicial, mas ainda reduzidos. Além disso, os tempos de resposta dependem do hardware local e do modelo em execução, pelo que os resultados de latência não devem ser generalizados sem indicar claramente as condições de teste.",
    ),
    (
        "p",
        "A quarta limitação respeita à qualidade dos dados de entrada. Como qualquer sistema de apoio à decisão em GRC, o Virtual CISO só é tão fiável quanto a completude e qualidade dos dados registados sobre ativos, vulnerabilidades, controlos, mecanismos e evidências. Organizações com baixa maturidade documental exigirão maior esforço inicial de parametrização.",
    ),
    ("h2", "Ameaças à validade"),
    (
        "p",
        "Em termos de validade de construção, existe o risco de alguns conceitos, como mecanismo, evidência esperada ou controlo interno, assumirem significados diferentes em organizações distintas. Esta ameaça foi mitigada pela modelação flexível, pela validação humana dos mappings e pela preservação de rationale em associações críticas.",
    ),
    (
        "p",
        "Quanto à validade interna, parte dos testes e da demonstração foi conduzida pelo próprio investigador, o que pode introduzir enviesamento na seleção de cenários e interpretação dos resultados. A utilização de comandos de avaliação, endpoints rastreáveis e placeholders para registo explícito dos resultados procura reduzir esta ameaça.",
    ),
    (
        "p",
        "No plano da validade externa, os resultados não podem ser automaticamente generalizados a organizações com outros setores, níveis de maturidade, frameworks ou infraestruturas. A modularidade da arquitetura sugere capacidade de adaptação, mas essa hipótese deverá ser validada em estudos futuros.",
    ),
    ("h2", "Síntese da demonstração e avaliação"),
    (
        "p",
        "A demonstração e avaliação apresentadas neste capítulo indicam que o Virtual CISO evoluiu para um artefacto coerente de apoio à decisão em GRC, centrado em controlos internos, mecanismos, evidências, documentos e rastreabilidade. O fluxo demonstrado permite ao CISO criar uma política, associar controlos internos, operacionalizar mecanismos, recolher evidências, mapear frameworks, calcular conformidade e consultar traceability.",
    ),
    (
        "p",
        "A avaliação da IA evidencia que o assistente é mais útil quando opera sobre contexto interno estruturado e indexado, em vez de responder apenas com conhecimento genérico. A abordagem internal-first, implementada por adapter sem alterar o router, reforça a coerência arquitetural e melhora a defensabilidade das respostas. A principal limitação observada é a latência dos modelos locais em ambiente com recursos reduzidos, aspeto que deve ser apresentado como compromisso entre privacidade, controlo e desempenho.",
    ),
    (
        "p",
        "Em síntese, os resultados sustentam a proposta do Virtual CISO como contributo original para o cruzamento entre plataformas de GRC, apoio à decisão em cibersegurança e IA explicável. Persistem limitações e trabalho futuro, mas a arquitetura desenvolvida oferece uma base sólida para validação com utilizadores, expansão dos testes, melhoria da performance e eventual adoção em contexto organizacional real.",
    ),
]


REFERENCES = [
    "Bass, L., Clements, P., & Kazman, R. (2021). Software Architecture in Practice (4.ª ed.). Addison-Wesley.",
    "European Parliament and Council. (2022). Directive (EU) 2022/2555 on measures for a high common level of cybersecurity across the Union (NIS2). Official Journal of the European Union.",
    "European Parliament and Council. (2024). Regulation (EU) 2024/1689 laying down harmonised rules on artificial intelligence. Official Journal of the European Union.",
    "FIRST. (2019). Common Vulnerability Scoring System v3.1: Specification Document. Forum of Incident Response and Security Teams.",
    "Hevner, A. R., March, S. T., Park, J., & Ram, S. (2004). Design science in information systems research. MIS Quarterly, 28(1), 75-105.",
    "ISO/IEC. (2022). ISO/IEC 27001:2022 Information security, cybersecurity and privacy protection - Information security management systems - Requirements.",
    "ISO/IEC. (2022). ISO/IEC 27002:2022 Information security, cybersecurity and privacy protection - Information security controls.",
    "Jacobs, J., Romanosky, S., Edwards, B., Adjerid, I., & Roytman, M. (2021). Exploit Prediction Scoring System (EPSS). Digital Threats: Research and Practice.",
    "Lewis, P., Perez, E., Piktus, A., Petroni, F., Karpukhin, V., Goyal, N., et al. (2020). Retrieval-Augmented Generation for Knowledge-Intensive NLP Tasks. Advances in Neural Information Processing Systems.",
    "Lundberg, S. M., & Lee, S.-I. (2017). A unified approach to interpreting model predictions. Advances in Neural Information Processing Systems.",
    "March, S. T., & Smith, G. F. (1995). Design and natural science research on information technology. Decision Support Systems, 15(4), 251-266.",
    "NIST. (2024). The NIST Cybersecurity Framework (CSF) 2.0. National Institute of Standards and Technology.",
    "Peffers, K., Tuunanen, T., Rothenberger, M. A., & Chatterjee, S. (2007). A design science research methodology for information systems research. Journal of Management Information Systems, 24(3), 45-77.",
    "Sommerville, I. (2016). Software Engineering (10.ª ed.). Pearson.",
]


def add_items(doc: Document, anchor: Paragraph, items: list[tuple[str, object]]) -> None:
    for kind, payload in items:
        if kind == "h1":
            add_page_break_before(doc, anchor)
            add_paragraph_before(doc, anchor, str(payload), "Heading 1")
        elif kind == "h2":
            add_paragraph_before(doc, anchor, str(payload), "Heading 2")
        elif kind == "h3":
            add_paragraph_before(doc, anchor, str(payload), "Heading 3")
        elif kind == "p":
            add_paragraph_before(doc, anchor, str(payload), "Normal")
        elif kind == "bullets":
            add_bullets(doc, anchor, list(payload))
        elif kind == "table":
            spec = dict(payload)
            add_table_before(doc, anchor, spec["caption"], spec["headers"], spec["rows"])
        else:
            raise ValueError(f"Unknown item kind: {kind}")


def insert_references(doc: Document) -> None:
    glossario = find_paragraph(doc, "Glossário")
    existing = "\n".join(p.text for p in doc.paragraphs)
    for ref in REFERENCES:
        if ref not in existing:
            add_paragraph_before(doc, glossario, ref, "Normal")


def main() -> None:
    doc = Document(SOURCE)
    anchor = remove_body_range(
        doc,
        "Conceção e Desenvolvimento do Virtual CISO",
        "Conclusões e Trabalho Futuro",
    )
    add_items(doc, anchor, CHAPTER_4)
    add_items(doc, anchor, CHAPTER_5)
    insert_references(doc)
    set_normal_font(doc)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
