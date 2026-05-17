from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT_PATH = r"D:\virtual_ciso\virtual_CISO_V1_2\docs\Capitulo_5_Avaliacao_Virtual_CISO.docx"


def set_run_font(run, name="Calibri", size=11, bold=False, color=None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def style_paragraph(paragraph, before=0, after=6, line=1.1):
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line


def add_heading(doc, text, level):
    p = doc.add_paragraph()
    run = p.add_run(text)
    if level == 1:
        set_run_font(run, size=16, bold=True, color="2E74B5")
        style_paragraph(p, before=16, after=8)
    elif level == 2:
        set_run_font(run, size=13, bold=True, color="2E74B5")
        style_paragraph(p, before=12, after=6)
    else:
        set_run_font(run, size=12, bold=True, color="1F4D78")
        style_paragraph(p, before=8, after=4)


def add_paragraph(doc, text):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_run_font(run, size=11)
    style_paragraph(p, before=0, after=6)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        run = p.add_run(item)
        set_run_font(run, size=11)
        style_paragraph(p, before=0, after=4, line=1.15)


def add_table(doc):
    doc.add_paragraph()
    table = doc.add_table(rows=1, cols=3)
    table.style = "Table Grid"
    table.autofit = False
    widths = [Inches(2.2), Inches(1.6), Inches(2.7)]
    hdr = table.rows[0].cells
    headers = ["Indicador", "Resultado", "Observação"]
    for idx, text in enumerate(headers):
        hdr[idx].width = widths[idx]
        p = hdr[idx].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(text)
        set_run_font(run, size=11, bold=True)

    rows = [
        ("LLMRouter - task_type", "100,00%", "Após refinamento das regras de precedência."),
        ("LLMRouter - needs_rag", "100,00%", "Decisão de ativação do RAG consistente no dataset avaliado."),
        ("Assistente - cenários estruturados", "100,00%", "2 em 2 cenários corretos, com latência muito baixa."),
        ("Assistente - pergunta geral", "100,00%", "Pergunta 'O que faz um CISO?' corretamente tratada como general_qa."),
        ("Assistente - priorização", "100,00%", "Fluxo funcional validado com fontes e resposta útil."),
    ]

    for indicator, result, obs in rows:
        cells = table.add_row().cells
        values = [indicator, result, obs]
        for idx, value in enumerate(values):
            cells[idx].width = widths[idx]
            p = cells[idx].paragraphs[0]
            if idx == 1:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(value)
            set_run_font(run, size=11)
    return table


def build():
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Capítulo 5 - Demonstração e Avaliação do Virtual CISO")
    set_run_font(run, size=18, bold=True, color="2E74B5")
    style_paragraph(title, before=0, after=14)

    add_heading(doc, "5.1. Enquadramento da avaliação", 2)
    add_paragraph(
        doc,
        "Após a conceção e desenvolvimento do artefacto, tornou-se necessário demonstrar e avaliar a sua capacidade para responder ao problema identificado na dissertação. No contexto da metodologia de Design Science Research, a avaliação não se limita à verificação técnica de componentes isoladas, procurando antes analisar em que medida o Virtual CISO cumpre os objetivos definidos, suporta os fluxos pretendidos e produz resultados úteis, rastreáveis e coerentes para o domínio de Governação, Risco e Conformidade em cibersegurança.",
    )
    add_paragraph(
        doc,
        "A avaliação foi conduzida em contexto de prova de conceito e incidiu sobre a plataforma como um todo. Assim, foram considerados os módulos de inventário, risco, conformidade, governação e apoio à decisão, bem como a camada de IA generativa e RAG híbrido que complementa esses módulos. Esta opção é coerente com a natureza do artefacto, já que o valor do sistema decorre sobretudo da integração entre dados estruturados, mecanismos analíticos, regras determinísticas, recuperação semântica e interface orientada à decisão.",
    )

    add_heading(doc, "5.2. Estratégia de demonstração e avaliação", 2)
    add_paragraph(
        doc,
        "A demonstração do artefacto foi realizada através de cenários representativos de utilização, construídos de forma a refletir necessidades reais de apoio ao CISO. Estes cenários incidiram sobre fluxos que articulam entidades técnicas, organizacionais e normativas, permitindo observar o comportamento integrado da plataforma em tarefas de inventário, análise de risco, priorização, conformidade e interação com o assistente em linguagem natural.",
    )
    add_bullets(
        doc,
        [
            "consulta e gestão do inventário de ativos;",
            "análise de vulnerabilidades e risco associado;",
            "priorização de intervenções com base em múltiplas dimensões;",
            "avaliação de conformidade por framework;",
            "associação entre controlos, mecanismos, evidência e findings;",
            "formulação de perguntas em linguagem natural ao assistente Virtual CISO.",
        ],
    )

    add_heading(doc, "5.3. Critérios de avaliação do artefacto", 2)
    add_paragraph(
        doc,
        "Para avaliar o Virtual CISO de forma consistente com os objetivos da dissertação, foram considerados quatro critérios principais: integração entre dimensões de GRC, apoio à decisão e priorização, explicabilidade e rastreabilidade, e interação em linguagem natural através da componente de IA híbrida.",
    )
    add_bullets(
        doc,
        [
            "capacidade de integrar informação dispersa de diferentes dimensões de GRC;",
            "capacidade de produzir resultados acionáveis para risco, conformidade e priorização;",
            "capacidade de justificar e rastrear recomendações e respostas;",
            "capacidade de suportar interação em linguagem natural através de uma componente de IA híbrida.",
        ],
    )

    add_heading(doc, "5.4. Demonstração funcional da plataforma", 2)
    add_paragraph(
        doc,
        "A avaliação funcional demonstrou que a plataforma já suporta uma base operacional coerente de inventário, permitindo gerir ativos, categorias, tipos, localizações, ambientes e infraestruturas. Esta capacidade é relevante porque o inventário constitui o ponto de partida para a correlação entre exposição técnica, criticidade organizacional e risco.",
    )
    add_paragraph(
        doc,
        "Foi igualmente possível observar a integração entre vulnerabilidades, métricas técnicas e informação contextual, suportando a priorização de intervenções. Em vez de depender exclusivamente da severidade técnica, a arquitetura implementada relaciona criticidade dos ativos, exposição, relevância normativa, mecanismos existentes e evidência disponível.",
    )
    add_paragraph(
        doc,
        "No domínio da governação e conformidade, a plataforma demonstrou uma lógica de conformidade contínua, associando frameworks, controlos, mecanismos, políticas, evidência e lacunas. Esta capacidade aproxima o artefacto de um modelo auditável e rastreável de apoio à conformidade.",
    )

    add_heading(doc, "5.5. Avaliação da componente de IA híbrida", 2)
    add_paragraph(
        doc,
        "A componente de IA do Virtual CISO foi avaliada como parte de um sistema de apoio à decisão mais amplo. O assistente não substitui os motores determinísticos nem funciona isoladamente do modelo de dados; atua como camada de interpretação, comunicação e síntese sobre informação previamente estruturada e recuperada.",
    )
    add_paragraph(
        doc,
        "Neste contexto, a avaliação incidiu sobre a integração entre consultas estruturadas por ORM, recuperação semântica de conhecimento em pgvector, priorização determinística de vulnerabilidades e construção de respostas fundamentadas com fontes.",
    )

    add_heading(doc, "5.5.1. Avaliação do LLMRouter", 3)
    add_paragraph(
        doc,
        "Para avaliar o LLMRouter, foi construído um conjunto rotulado de perguntas representativas dos principais fluxos funcionais da plataforma. A avaliação foi automatizada através de um comando dedicado, permitindo medir a correção do task_type previsto, a decisão de ativação do RAG e a fonte de decisão utilizada pelo router.",
    )
    add_paragraph(
        doc,
        "Numa fase inicial, o router obteve 90,48% de acurácia em task_type e 95,24% em needs_rag, tendo sido identificados dois erros relevantes: uma pergunta de priorização de vulnerabilidades confundida com structured_query e uma pergunta conceptual sobre o papel do CISO confundida com executive_advisory. Após o refinamento das regras de precedência e a introdução de regras explícitas para perguntas conceptuais gerais, a avaliação passou a 100,00% em task_type, 100,00% em needs_rag e 100,00% de acerto exato no dataset avaliado.",
    )

    add_heading(doc, "5.5.2. Avaliação funcional do assistente", 3)
    add_paragraph(
        doc,
        "A avaliação funcional do assistente incidiu sobre o pipeline completo, incluindo routing, structured query, recuperação RAG, geração de resposta e fontes devolvidas. Para esse efeito, foi criado um dataset próprio de cenários, permitindo observar tanto o comportamento correto da classificação como a utilidade prática das respostas produzidas.",
    )
    add_paragraph(
        doc,
        "Nos cenários estruturados, o assistente obteve 100,00% de acerto funcional, recorrendo ao bypass ORM e apresentando latências muito baixas, na ordem de dezenas a poucas centenas de milissegundos. Nos cenários generativos, a avaliação revelou inicialmente um gargalo de latência associado ao serviço LLM remoto, com tempos superiores a um minuto e falhas por timeout em casos mais pesados, como a priorização de vulnerabilidades.",
    )
    add_paragraph(
        doc,
        "Após o aumento controlado do timeout, a introdução de fallback determinístico para a priorização de vulnerabilidades e o ajuste do perfil de geração por tarefa, observou-se uma melhoria clara da latência sem perda de correção funcional. A pergunta geral 'O que faz um CISO?' passou a apresentar uma latência aproximada de 22 segundos, quando antes se encontrava próxima de 96 segundos. No caso da pergunta 'Prioriza as vulnerabilidades mais urgentes', a latência desceu de cerca de 174 segundos para aproximadamente 86 segundos, mantendo resposta útil e fontes consistentes.",
    )

    add_heading(doc, "5.5.3. Síntese quantitativa", 3)
    add_paragraph(
        doc,
        "A Tabela 5.1 sintetiza os principais resultados quantitativos observados na avaliação do router e do assistente.",
    )
    add_table(doc)
    add_paragraph(doc, "Tabela 5.1 - Síntese dos resultados da avaliação do LLMRouter e do assistente.")

    add_heading(doc, "5.6. Discussão global da avaliação", 2)
    add_paragraph(
        doc,
        "Os resultados observados permitem concluir que o Virtual CISO já demonstra viabilidade como artefacto de apoio à decisão, na medida em que integra múltiplas dimensões de GRC, disponibiliza fluxos funcionais coerentes e incorpora mecanismos de apoio à interpretação, rastreabilidade e recomendação. Do ponto de vista funcional, a plataforma evidencia capacidade para representar e correlacionar ativos, vulnerabilidades, risco, controlos, mecanismos, evidência e conformidade. Do ponto de vista da IA, os resultados mostram que uma abordagem híbrida, combinando regras, ORM, RAG e LLM, é mais adequada do que uma abordagem puramente generativa para o domínio em análise.",
    )

    add_heading(doc, "5.7. Limitações da avaliação", 2)
    add_paragraph(
        doc,
        "A avaliação realizada apresenta algumas limitações que importa explicitar. Em primeiro lugar, o número de queries usadas na avaliação do LLMRouter e do assistente é ainda reduzido, o que limita a generalização estatística dos resultados. Em segundo lugar, a avaliação incidiu sobretudo sobre cenários representativos construídos para prova de conceito, não correspondendo a um estudo longitudinal em ambiente real.",
    )
    add_paragraph(
        doc,
        "Em terceiro lugar, a qualidade das respostas do assistente depende da qualidade e cobertura dos dados existentes no sistema, bem como do estado da base vetorial de conhecimento e da disponibilidade do serviço LLM. Por essa razão, o desempenho da componente de IA deve ser entendido como dependente da maturidade global do artefacto e não apenas do modelo de linguagem utilizado.",
    )

    add_heading(doc, "5.8. Síntese do capítulo", 2)
    add_paragraph(
        doc,
        "O presente capítulo demonstrou e avaliou o Virtual CISO enquanto artefacto de apoio à decisão para CISOs, adotando uma perspetiva abrangente que incluiu módulos funcionais, mecanismos de integração, capacidade de priorização, rastreabilidade e componente de IA híbrida. Os resultados obtidos indicam que a plataforma já materializa, em contexto de prova de conceito, uma resposta consistente às lacunas identificadas na revisão sistemática, nomeadamente a fragmentação entre GRC, a dificuldade de correlacionar risco com conformidade e a necessidade de maior explicabilidade em sistemas baseados em IA.",
    )
    add_paragraph(
        doc,
        "No caso específico da componente de IA, a avaliação do LLMRouter e do assistente mostrou resultados encorajadores, com acurácia elevada, melhoria mensurável após refinamento e redução importante da latência após otimização do perfil de geração. Esta análise reforça a ideia de que a avaliação do Virtual CISO deve ser entendida como avaliação de um artefacto integrado, no qual a IA desempenha um papel relevante, mas sempre articulado com dados estruturados, lógica determinística e validação humana.",
    )

    doc.save(OUTPUT_PATH)
    print(OUTPUT_PATH)


if __name__ == "__main__":
    build()
